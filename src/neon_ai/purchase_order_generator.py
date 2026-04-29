import os

from docx import Document


def _safe_filename(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in (" ", "-", "_") else "_" for ch in str(value or ""))
    cleaned = "_".join(cleaned.split())
    return cleaned.strip("_") or "Purchase_Order"


def generate_po_docx(header_data, items, save_path=None):
    """Builds the PO DOCX and saves it to the requested project-folder path."""
    doc = Document()

    po_id = header_data["PurchaseOrderID"]
    vendor_name = header_data.get("VendorName") or "Vendor"
    customer_name = header_data.get("CustomerName") or "Customer"
    site_name = header_data.get("SiteName") or "Site"
    street_number = header_data.get("StreetNumber") or ""
    street_name = header_data.get("StreetName") or ""
    site_address = " ".join(part for part in (street_number, street_name) if part).strip() or site_name

    doc.add_heading(f"PURCHASE ORDER #{po_id}", 0)

    intro = doc.add_paragraph()
    intro.add_run("Vendor: ").bold = True
    intro.add_run(vendor_name)
    intro.add_run("\nVendor Account #: ").bold = True
    intro.add_run(str(header_data.get("AccountNumber") or "N/A"))
    intro.add_run("\nProject: ").bold = True
    intro.add_run(customer_name)
    intro.add_run("\nShip To: ").bold = True
    intro.add_run(site_name)
    intro.add_run(f"\nSite Address: {site_address}")
    intro.add_run(f"\nPO Date: {header_data.get('Date') or 'N/A'}")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Qty"
    hdr_cells[1].text = "Description"
    hdr_cells[2].text = "Unit Price"
    hdr_cells[3].text = "Ext Price"

    grand_total = 0.0
    for item in items:
        qty = float(item.get("qty") or 0)
        price = float(item.get("price") or 0)
        ext_price = qty * price
        row_cells = table.add_row().cells
        row_cells[0].text = f"{qty:.2f}".rstrip("0").rstrip(".")
        row_cells[1].text = str(item.get("desc") or "")
        row_cells[2].text = f"${price:,.2f}"
        row_cells[3].text = f"${ext_price:,.2f}"
        grand_total += ext_price

    total_paragraph = doc.add_paragraph()
    total_paragraph.add_run("TOTAL PURCHASE VALUE: ").bold = True
    total_paragraph.add_run(f"${grand_total:,.2f}")

    if not save_path:
        filename = f"PO_{po_id}_{_safe_filename(vendor_name)}.docx"
        save_path = os.path.join(os.getcwd(), filename)

    os.makedirs(os.path.dirname(save_path), exist_ok=True)
    doc.save(save_path)
    print(f"PO #{po_id} DOCX saved to: {save_path}")
    return save_path


def convert_docx_to_pdf(docx_path, pdf_path=None):
    """Converts a DOCX to PDF using local Microsoft Word automation on Windows."""
    if not docx_path or not os.path.exists(docx_path):
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    if not pdf_path:
        pdf_path = os.path.splitext(docx_path)[0] + ".pdf"

    word = None
    document = None
    try:
        from win32com import client  # type: ignore

        word = client.DispatchEx("Word.Application")
        word.Visible = False
        document = word.Documents.Open(os.path.abspath(docx_path))
        document.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
        return pdf_path
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
