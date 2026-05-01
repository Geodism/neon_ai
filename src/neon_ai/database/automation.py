import os
import re
import datetime
import textwrap
from typing import Dict, List, Optional
from docx import Document
from neon_ai.database.connection import get_connection
from neon_ai.database.folders import get_target_folder
from neon_ai.doc_generator import export_estimate_doc
from neon_ai.database.estimates import get_detailed_estimate_data, ensure_estimate_sequence
from neon_ai.database.customers import sync_table_sequence

# ==========================================
# 1. THE GOLDEN RULES (LIBRARIAN & AUDIT TRAIL)
# ==========================================
MASTER_DRIVE = "D:/Projects"
# Automation is now fully ungated; keep a very early cutoff so the existing
# sweeper queries still work without excluding historical records.
AUTOMATION_MIN_DATE = datetime.date(1900, 1, 1)
LEGACY_ESTIMATE_AUTO_SEND_ENV = "NEON_ENABLE_LEGACY_ESTIMATE_AUTO_SEND"

def clean_name(name):
    """Removes illegal Windows characters."""
    if not name: return ""
    return re.sub(r'[\\/*?:"<>|]', "", str(name)).strip()

def get_project_file_paths(customer_name, site_address, estimate_id, category=None):
    """Handles the D: to C: failover and builds consistent folder paths."""
    safe_cust = clean_name(customer_name)
    safe_site = clean_name(site_address)
    
    project_folder = f"{safe_site} - {safe_cust}".strip()
    
    if category:
        target_path = os.path.join(MASTER_DRIVE, project_folder, str(estimate_id), category)
    else:
        target_path = os.path.join(MASTER_DRIVE, project_folder, str(estimate_id))

    target_path = target_path.replace("\\", "/")

    drive = os.path.splitdrive(target_path)[0]
    if not os.path.exists(drive):
        target_path = target_path.replace("D:", "C:")
    
    os.makedirs(target_path, exist_ok=True)
    docx_path = os.path.join(target_path, f"Lead_Sheet_{estimate_id}.docx")
    
    return target_path, docx_path

def log_estimate_action(estimate_id: int, note_text: str, category: str = 'AI Action'):
    """The Audit Trail: Gives the AI 'memory' by logging actions to the database."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute('''
            INSERT INTO "Note" ("EstimateID", "NoteText", "Category") 
            VALUES (%s, %s, %s)
        ''', (estimate_id, note_text, category))
        conn.commit()
        print(f"📝 Audit Logged: {note_text}")
        return True
    except Exception as e:
        print(f"Audit Log Error: {e}")
        conn.rollback()
        return False
    finally:
        conn.close()


def _verification_value_missing(value) -> bool:
    text = str(value or "").strip().lower()
    return text in {"", "pending", "null", "none", "unknown", "n/a"}


def get_estimate_verification_snapshot(estimate_id: int) -> Optional[Dict[str, object]]:
    """Captures the key estimate/customer/site fields we care about for action verification."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            SELECT
                e."EstimateID",
                e."Status" AS "EstimateStatus",
                e."Description",
                c."CustomerID",
                c."CustomerName",
                c."Email",
                c."Phone",
                c."Address",
                c."StreetName" AS "CustomerStreetName",
                c."CityName" AS "CustomerCityName",
                c."PostalCode",
                s."SiteID",
                s."SiteName",
                s."StreetNumber" AS "SiteStreetNumber",
                s."StreetName" AS "SiteStreetName",
                s."City" AS "SiteCity"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."EstimateID" = %s
            ''',
            (estimate_id,),
        )
        row = cur.fetchone()
        if not row:
            return None

        missing_fields = []
        if _verification_value_missing(row.get("CustomerName")):
            missing_fields.append("CustomerName")
        if _verification_value_missing(row.get("Email")):
            missing_fields.append("CustomerEmail")
        if _verification_value_missing(row.get("Phone")):
            missing_fields.append("CustomerPhone")
        if _verification_value_missing(row.get("Address")):
            missing_fields.append("AddressNumber")
        if _verification_value_missing(row.get("CustomerStreetName")):
            missing_fields.append("StreetName")
        if _verification_value_missing(row.get("CustomerCityName")):
            missing_fields.append("CityName")
        if _verification_value_missing(row.get("PostalCode")):
            missing_fields.append("PostalCode")
        if _verification_value_missing(row.get("SiteName")) or "pending site" in str(row.get("SiteName") or "").lower():
            missing_fields.append("SiteAddress")

        snapshot = dict(row)
        snapshot["MissingFields"] = missing_fields
        return snapshot
    finally:
        conn.close()


def _flatten_verification_snapshot(snapshot: Optional[Dict[str, object]]) -> Dict[str, object]:
    if not snapshot:
        return {}
    return {
        "EstimateStatus": snapshot.get("EstimateStatus"),
        "CustomerName": snapshot.get("CustomerName"),
        "CustomerEmail": snapshot.get("Email"),
        "CustomerPhone": snapshot.get("Phone"),
        "AddressNumber": snapshot.get("Address"),
        "StreetName": snapshot.get("CustomerStreetName"),
        "CityName": snapshot.get("CustomerCityName"),
        "PostalCode": snapshot.get("PostalCode"),
        "SiteName": snapshot.get("SiteName"),
        "SiteStreetNumber": snapshot.get("SiteStreetNumber"),
        "SiteStreetName": snapshot.get("SiteStreetName"),
        "SiteCity": snapshot.get("SiteCity"),
        "MissingFields": tuple(snapshot.get("MissingFields", [])),
    }


def _format_snapshot_summary(snapshot: Optional[Dict[str, object]]) -> str:
    if not snapshot:
        return "no record"
    missing = ", ".join(snapshot.get("MissingFields", [])) or "None"
    return (
        f"status={snapshot.get('EstimateStatus') or 'None'}; "
        f"customer={snapshot.get('CustomerName') or 'None'}; "
        f"email={snapshot.get('Email') or 'None'}; "
        f"phone={snapshot.get('Phone') or 'None'}; "
        f"site={snapshot.get('SiteName') or 'None'}; "
        f"missing={missing}"
    )


def _format_verification_value(value) -> str:
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item) for item in value) or "None"
    return str(value if value not in (None, "") else "None")


def log_estimate_verification(
    estimate_id: int,
    action_label: str,
    before_snapshot: Optional[Dict[str, object]],
    after_snapshot: Optional[Dict[str, object]],
    expected_values: Optional[Dict[str, object]] = None,
    expected_changed_fields: Optional[List[str]] = None,
):
    """Writes a verification note that compares before/after database state for an action."""
    before_flat = _flatten_verification_snapshot(before_snapshot)
    after_flat = _flatten_verification_snapshot(after_snapshot)

    observed_changes = []
    for field_name in sorted(set(before_flat.keys()) | set(after_flat.keys())):
        before_value = before_flat.get(field_name)
        after_value = after_flat.get(field_name)
        if before_value != after_value:
            observed_changes.append(
                f"{field_name}: '{_format_verification_value(before_value)}' -> '{_format_verification_value(after_value)}'"
            )

    checks = []
    passed = True

    for field_name, expected_value in (expected_values or {}).items():
        actual_value = after_flat.get(field_name)
        ok = actual_value == expected_value
        checks.append(
            f"{field_name} expected '{_format_verification_value(expected_value)}' got "
            f"'{_format_verification_value(actual_value)}' [{'PASS' if ok else 'FAIL'}]"
        )
        passed = passed and ok

    for field_name in (expected_changed_fields or []):
        changed = before_flat.get(field_name) != after_flat.get(field_name)
        checks.append(f"{field_name} changed [{'PASS' if changed else 'FAIL'}]")
        passed = passed and changed

    if not checks:
        checks.append("No explicit verification checks supplied [INFO]")

    note_text = (
        f"Verification {'PASS' if passed else 'FAIL'} for {action_label}. "
        f"Before: {_format_snapshot_summary(before_snapshot)}. "
        f"After: {_format_snapshot_summary(after_snapshot)}. "
        f"Observed DB changes: {'; '.join(observed_changes) if observed_changes else 'None'}. "
        f"Checks: {'; '.join(checks)}."
    )
    log_estimate_action(estimate_id, note_text, "Verification")
    return passed


def get_customer_intake_status(estimate_id: int):
    """Returns the linked customer and any still-missing intake fields for an estimate."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            SELECT
                c."CustomerID",
                c."CustomerName",
                c."Email",
                c."Phone",
                c."Address",
                c."StreetName",
                c."CityName",
                c."PostalCode"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."EstimateID" = %s
            ''',
            (estimate_id,),
        )
        row = cur.fetchone()
        if not row:
            return None

        def _missing(value):
            text = str(value or "").strip().lower()
            return text in {"", "pending", "null", "none", "unknown", "n/a"}

        missing = []
        if _missing(row["CustomerName"]):
            missing.append("CustomerName")
        if _missing(row["Email"]):
            missing.append("CustomerEmail")
        if _missing(row["Phone"]):
            missing.append("CustomerPhone")
        if _missing(row["Address"]):
            missing.append("AddressNumber")
        if _missing(row["StreetName"]):
            missing.append("StreetName")
        if _missing(row["CityName"]):
            missing.append("CityName")
        if _missing(row["PostalCode"]):
            missing.append("PostalCode")

        return {
            "customer_id": row["CustomerID"],
            "customer_name": row["CustomerName"],
            "email": row["Email"],
            "missing": missing,
        }
    finally:
        conn.close()


def has_customer_info_request_been_sent(estimate_id: int, missing_fields=None):
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            SELECT "NoteText"
            FROM "Note"
            WHERE "EstimateID" = %s
              AND "Category" = 'Customer Intake'
            ORDER BY "Timestamp" DESC
            ''',
            (estimate_id,),
        )
        rows = cur.fetchall()
    finally:
        conn.close()

    if not rows:
        return False

    if not missing_fields:
        return True

    expected = {field.strip() for field in missing_fields}
    for row in rows:
        note_text = str(row.get("NoteText") or "")
        if "customer info request sent" not in note_text.lower():
            continue
        matched = {field for field in expected if field.lower() in note_text.lower()}
        if matched == expected:
            return True
    return False


def format_customer_missing_fields(missing_fields):
    friendly = {
        "CustomerName": "your full name",
        "CustomerEmail": "the best email address to use",
        "CustomerPhone": "your phone number",
        "AddressNumber": "your street number",
        "StreetName": "your street name",
        "CityName": "your city",
        "PostalCode": "your postal code",
    }
    return [friendly.get(field, field) for field in missing_fields]


def has_pending_customer_verification(estimate_id: int):
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            SELECT "NoteText"
            FROM "Note"
            WHERE "EstimateID" = %s
              AND "Category" = 'Customer Verification'
            ORDER BY "Timestamp" DESC
            ''',
            (estimate_id,),
        )
        rows = cur.fetchall()
    finally:
        conn.close()

    if not rows:
        return False

    latest_request_index = None
    for index, row in enumerate(rows):
        note_text = str(row.get("NoteText") or "").lower()
        if "verification request sent" in note_text:
            latest_request_index = index
            break

    if latest_request_index is None:
        return False

    for row in rows[:latest_request_index]:
        note_text = str(row.get("NoteText") or "").lower()
        if "verification reply processed" in note_text:
            return False
    return True


def send_customer_address_verification_request(customer_id: int):
    """Send a customer-facing verification email for contact + site address details."""
    from neon_ai.gateway import send_to_user
    from neon_ai.database.customers import (
        get_customer_by_id,
        get_customer_primary_site,
        get_customer_latest_estimate_id,
        get_customer_missing_fields,
    )

    customer = get_customer_by_id(customer_id)
    if not customer:
        raise ValueError(f"Customer #{customer_id} could not be found.")

    email_address = customer.get("Email")
    if not email_address:
        raise ValueError("This customer does not have an email address on file yet.")

    primary_site = get_customer_primary_site(customer_id)
    estimate_id = get_customer_latest_estimate_id(customer_id)
    missing_fields = get_customer_missing_fields(customer_id)
    contact_address = " ".join(
        part for part in [
            customer.get("Address") or "",
            customer.get("StreetName") or "",
            customer.get("CityName") or "",
            customer.get("PostalCode") or "",
        ]
        if part
    ) or "Not on file"

    site_address = " ".join(
        part for part in [
            primary_site.get("StreetNumber") if primary_site else "",
            primary_site.get("StreetName") if primary_site else "",
            primary_site.get("City") if primary_site else "",
        ]
        if part
    ) or (primary_site.get("SiteName") if primary_site else "Not on file")

    missing_text = ""
    if missing_fields:
        friendly_missing = "\n".join(f"- {field}" for field in format_customer_missing_fields(missing_fields))
        missing_text = (
            "\nWe are also missing the following pieces of contact information, so please include them in your reply:\n"
            f"{friendly_missing}\n"
        )

    subject = f"Please verify your contact and project address details"
    body = (
        f"Hello {customer.get('CustomerName') or 'there'},\n\n"
        "We are tightening up our project records and would appreciate a quick address check.\n\n"
        "Here is the customer contact information we currently have on file:\n"
        f"- Name: {customer.get('CustomerName') or 'Not on file'}\n"
        f"- Email: {customer.get('Email') or 'Not on file'}\n"
        f"- Phone: {customer.get('Phone') or 'Not on file'}\n"
        f"- Contact Address: {contact_address}\n\n"
        "Here is the project site information we currently have:\n"
        f"- Site: {primary_site.get('SiteName') if primary_site else 'Not on file'}\n"
        f"- Site Address: {site_address}\n\n"
        "Please reply with any corrections and let us know:\n"
        "1. Is the project site address the same as the customer contact address?\n"
        "2. If not, what is the correct project site address?\n"
        f"{missing_text}\n"
        "Thank you,\nArgon Electrical"
    )

    if not send_to_user(subject, body, recipient=email_address):
        raise RuntimeError("SMTP failed while sending the verification request.")

    if estimate_id:
        log_estimate_action(
            estimate_id,
            f"Customer verification request sent for customer #{customer_id}.",
            "Customer Verification",
        )

    return {
        "customer_id": customer_id,
        "estimate_id": estimate_id,
        "email": email_address,
        "site_id": primary_site.get("SiteID") if primary_site else None,
    }

def build_client_estimate_doc(estimate_id: int):
    """Builds the canonical customer-facing estimate document in the project folder."""
    data = get_detailed_estimate_data(estimate_id)
    if not data or not data.get("parent"):
        raise ValueError(f"Estimate #{estimate_id} could not be loaded.")

    p = data["parent"]
    lab_mu = float(p.get('LaborMarkUp') or 20.0) / 100
    mat_mu = float(p.get('MaterialMarkUp') or 20.0) / 100

    packet = {
        "id": p["EstimateID"],
        "date": p["CreatedDate"],
        "customer_name": p["CustomerName"],
        "customer_info": (
            f"{p.get('CustAddr') or ''} {p.get('CustStreet') or ''}\n"
            f"{p.get('CustCity') or ''}, {p.get('CustZip') or ''}\n"
            f"Ph: {p.get('Phone') or 'N/A'}\n"
            f"Email: {p.get('Email') or 'N/A'}"
        ).strip(),
        "site_name": p["SiteName"],
        "site_addr": f"{p.get('SiteNum') or ''} {p.get('SiteStreet') or ''}, {p.get('SiteCity') or ''}".strip(" ,"),
        "scope": p.get("Description") or "Standard electrical scope.",
        "labor_items": [
            {
                "role": l["RoleDescription"],
                "total": (float(l["Hours"]) * float(l["Rate"])) * (1 + lab_mu)
            }
            for l in data["labor"]
        ],
        "mat_items": [
            {
                "desc": m["Description"],
                "total": (float(m["Quantity"]) * float(m["UnitCost"])) * (1 + mat_mu)
            }
            for m in data["materials"]
        ]
    }
    target_dir = get_target_folder(
        p.get("SiteNum"),
        p.get("SiteStreet"),
        p.get("SiteName"),
        p["EstimateID"],
        "Estimates"
    )
    filename = f"Customer_Estimate_{p['EstimateID']}.docx"
    return export_estimate_doc(packet, destination_dir=target_dir, filename=filename)

def get_client_estimate_total(estimate_id: int):
    """Calculates the same sell total used in the customer-facing estimate document."""
    data = get_detailed_estimate_data(estimate_id)
    if not data or not data.get("parent"):
        raise ValueError(f"Estimate #{estimate_id} could not be loaded.")

    p = data["parent"]
    lab_mu = float(p.get("LaborMarkUp") or 20.0) / 100
    mat_mu = float(p.get("MaterialMarkUp") or 20.0) / 100

    labor_total = sum((float(l["Hours"]) * float(l["Rate"])) * (1 + lab_mu) for l in data["labor"])
    material_total = sum((float(m["Quantity"]) * float(m["UnitCost"])) * (1 + mat_mu) for m in data["materials"])
    return labor_total + material_total

def get_client_estimate_doc_path(estimate_id: int):
    """Returns the canonical customer-facing estimate document path without regenerating it."""
    data = get_detailed_estimate_data(estimate_id)
    if not data or not data.get("parent"):
        raise ValueError(f"Estimate #{estimate_id} could not be loaded.")

    p = data["parent"]
    target_dir = get_target_folder(
        p.get("SiteNum"),
        p.get("SiteStreet"),
        p.get("SiteName"),
        p["EstimateID"],
        "Estimates"
    )
    return os.path.join(target_dir, f"Customer_Estimate_{p['EstimateID']}.docx")

def record_estimate_workflow_handoff(estimate_id, route_to_step, proof_type=None, proof_reference=None):
    """Leaves an explicit workflow breadcrumb for the human PM before conversion."""
    note_bits = [f"Workflow handoff prepared for step: {route_to_step}."]
    if proof_type:
        note_bits.append(f"AcceptanceProofType={proof_type}.")
    if proof_reference:
        note_bits.append(f"AcceptanceProofFileID={proof_reference}.")
    log_estimate_action(estimate_id, " ".join(note_bits), "Workflow")

def find_estimate_by_customer_email(customer_email, statuses=None):
    """Finds the newest estimate for a customer email in one of the given statuses."""
    statuses = statuses or ["Sent"]
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT e."EstimateID", e."Status", e."Description", s."SiteName", c."CustomerName", c."Email"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE LOWER(TRIM(c."Email")) = LOWER(TRIM(%s))
              AND e."Status" = ANY(%s)
            ORDER BY e."EstimateID" DESC
            LIMIT 1
        """, (customer_email, statuses))
        return cur.fetchone()
    finally:
        conn.close()

def get_sent_estimates_for_customer(customer_email):
    """Returns all sent estimates for a customer email, newest first."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT e."EstimateID", e."Status", e."Description", s."SiteName", c."CustomerName", c."Email"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE LOWER(TRIM(c."Email")) = LOWER(TRIM(%s))
              AND e."Status" = 'Sent'
            ORDER BY e."EstimateID" DESC
        """, (customer_email,))
        return cur.fetchall()
    finally:
        conn.close()

def get_estimates_for_customer_by_status(customer_email, statuses):
    """Returns customer estimates filtered by one or more statuses, newest first."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT e."EstimateID", e."Status", e."Description", s."SiteName", c."CustomerName", c."Email"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE LOWER(TRIM(c."Email")) = LOWER(TRIM(%s))
              AND e."Status" = ANY(%s)
            ORDER BY e."EstimateID" DESC
        """, (customer_email, statuses))
        return cur.fetchall()
    finally:
        conn.close()

def extract_estimate_id_from_text(email_subject, email_body):
    """Extracts an estimate number from email text like 'Estimate #123'."""
    text = f"{email_subject}\n{email_body}"
    patterns = [
        r'estimate\s*#\s*(\d+)',
        r'estimate\s+number\s*[:#]?\s*(\d+)',
        r'est\.?\s*#\s*(\d+)'
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None

def get_estimate_contact_context(estimate_id):
    """Fetches customer/site context for an estimate so file and email workflows share one lookup."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT
                e."EstimateID",
                e."Description",
                e."Status",
                e."SubmitDate",
                c."CustomerName",
                c."Email",
                s."SiteName",
                s."StreetNumber",
                s."StreetName"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."EstimateID" = %s
        """, (estimate_id,))
        return cur.fetchone()
    finally:
        conn.close()

def _pdf_escape(value):
    return (
        str(value or "")
        .replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )

def save_email_as_simple_pdf(pdf_path, title, lines):
    """Creates a simple text PDF without external dependencies."""
    page_width = 612
    page_height = 792
    left_margin = 54
    top_start = 738
    line_height = 14
    bottom_margin = 54

    wrapped_lines = []
    for raw_line in lines:
        normalized = str(raw_line or "")
        chunks = textwrap.wrap(normalized, width=92) or [""]
        wrapped_lines.extend(chunks)

    pages = []
    current_page = []
    y = top_start

    title_lines = textwrap.wrap(title or "Customer Response", width=80) or ["Customer Response"]
    for title_line in title_lines:
        current_page.append((left_margin, y, title_line))
        y -= line_height
    y -= 8

    for line in wrapped_lines:
        if y <= bottom_margin:
            pages.append(current_page)
            current_page = []
            y = top_start
        current_page.append((left_margin, y, line))
        y -= line_height

    if not current_page:
        current_page = [(left_margin, top_start, "")]
    pages.append(current_page)

    objects = []
    objects.append("<< /Type /Catalog /Pages 2 0 R >>")

    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Count {len(pages)} /Kids [{kids}] >>")

    for i, page_lines in enumerate(pages):
        page_obj_num = 3 + i * 2
        content_obj_num = page_obj_num + 1
        page_obj = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_width} {page_height}] "
            f"/Resources << /Font << /F1 {3 + len(pages) * 2} 0 R >> >> /Contents {content_obj_num} 0 R >>"
        )
        objects.append(page_obj)

        content_lines = ["BT", "/F1 11 Tf"]
        for x, y_pos, text in page_lines:
            content_lines.append(f"1 0 0 1 {x} {y_pos} Tm ({_pdf_escape(text)}) Tj")
        content_lines.append("ET")
        content_stream = "\n".join(content_lines)
        objects.append(f"<< /Length {len(content_stream.encode('latin-1', errors='replace'))} >>\nstream\n{content_stream}\nendstream")

    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    pdf_parts = ["%PDF-1.4\n"]
    offsets = [0]
    running_length = len(pdf_parts[0].encode("latin-1"))

    for index, obj in enumerate(objects, start=1):
        offsets.append(running_length)
        serialized = f"{index} 0 obj\n{obj}\nendobj\n"
        pdf_parts.append(serialized)
        running_length += len(serialized.encode("latin-1", errors="replace"))

    xref_offset = running_length
    xref_lines = [f"xref\n0 {len(objects) + 1}\n", "0000000000 65535 f \n"]
    for offset in offsets[1:]:
        xref_lines.append(f"{offset:010d} 00000 n \n")

    trailer = (
        "".join(xref_lines) +
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    )
    pdf_parts.append(trailer)

    with open(pdf_path, "wb") as handle:
        handle.write("".join(pdf_parts).encode("latin-1", errors="replace"))

    return pdf_path

def archive_estimate_response_pdf(estimate_id, customer_email, email_subject, email_body, message_id=None, response_type="Accepted"):
    """Saves a customer response as a PDF in the project folder."""
    context = get_estimate_contact_context(estimate_id)
    if not context:
        raise ValueError(f"Estimate #{estimate_id} could not be loaded for PDF archive.")

    target_dir = get_target_folder(
        context.get("StreetNumber"),
        context.get("StreetName"),
        context.get("SiteName"),
        context["EstimateID"],
        "Customer Responses"
    )
    safe_response = re.sub(r"[^A-Za-z0-9_-]+", "_", response_type or "Response").strip("_") or "Response"
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    pdf_path = os.path.join(target_dir, f"{safe_response}_Email_Estimate_{estimate_id}_{timestamp}.pdf")

    lines = [
        f"Estimate ID: {estimate_id}",
        f"Response Type: {response_type}",
        f"Customer: {context.get('CustomerName') or 'Unknown'}",
        f"Customer Email: {customer_email}",
        f"Site: {context.get('SiteName') or 'Unknown'}",
        f"Message-ID: {message_id or 'Not Provided'}",
        "",
        "Subject:",
        email_subject or "",
        "",
        "Body:",
        email_body or ""
    ]
    save_email_as_simple_pdf(
        pdf_path,
        f"Estimate #{estimate_id} {response_type} Email Archive",
        lines
    )
    return pdf_path

def has_follow_up_reminder_been_sent(estimate_id):
    """Checks whether the 48-hour follow-up reminder has already been sent to the owner."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT 1
            FROM "Note"
            WHERE "EstimateID" = %s
              AND "Category" = 'Follow Up'
              AND "NoteText" ILIKE '48-hour follow-up reminder sent%%'
            LIMIT 1
        """, (estimate_id,))
        return cur.fetchone() is not None
    finally:
        conn.close()

def sweep_for_sent_estimate_followups():
    """Notifies the owner when a sent estimate has been waiting 48+ hours without a customer decision."""
    from neon_ai.gateway import send_to_user

    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT
                e."EstimateID",
                e."Description",
                e."SubmitDate",
                c."CustomerName",
                c."Email",
                s."SiteName"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."Status" = 'Sent'
              AND e."SubmitDate" IS NOT NULL
              AND e."SubmitDate" >= %s
              AND e."SubmitDate" <= CURRENT_DATE - 2
            ORDER BY e."SubmitDate" ASC, e."EstimateID" ASC
        """, (AUTOMATION_MIN_DATE,))
        aging_estimates = cur.fetchall()

        if not aging_estimates:
            print("Sweeper: No 48-hour estimate follow-ups needed.")
            return

        for est in aging_estimates:
            if has_follow_up_reminder_been_sent(est["EstimateID"]):
                continue

            send_to_user(
                subject=f"Follow Up Needed: Estimate #{est['EstimateID']} for {est['CustomerName']}",
                content=(
                    f"Estimate #{est['EstimateID']} has been in Sent status for at least 48 hours without an acceptance or rejection.\n\n"
                    f"Customer: {est['CustomerName']}\n"
                    f"Email: {est['Email']}\n"
                    f"Site: {est['SiteName']}\n"
                    f"SubmitDate: {est['SubmitDate']}\n"
                    f"Scope: {est['Description']}\n\n"
                    "Recommended next step: follow up with the customer and offer any additional information they need to make a decision."
                )
            )
            log_estimate_action(
                est["EstimateID"],
                f"48-hour follow-up reminder sent to owner for Estimate #{est['EstimateID']}.",
                "Follow Up"
            )
    finally:
        conn.close()

def has_exported_invoice_for_estimate(estimate_id):
    """Checks whether an accepted estimate has any exported customer invoice through its work order."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT 1
            FROM "WorkOrder" wo
            JOIN "Invoice" i ON i."WorkOrderID" = wo."WorkOrderID"
            WHERE wo."SourceEstimateID" = %s
              AND i."InvoiceStatus" = 'Exported'
            LIMIT 1
        """, (estimate_id,))
        return cur.fetchone() is not None
    finally:
        conn.close()

def has_deposit_invoice_reminder_been_sent(estimate_id):
    """Checks whether the owner has already been reminded about the missing deposit invoice."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT 1
            FROM "Note"
            WHERE "EstimateID" = %s
              AND "Category" = 'Billing'
              AND "NoteText" ILIKE '24-hour deposit invoice reminder sent%%'
            LIMIT 1
        """, (estimate_id,))
        return cur.fetchone() is not None
    finally:
        conn.close()

def has_rfq_completion_nudge_been_sent(estimate_id):
    """Checks whether the owner was already nudged after every RFQ was resolved."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT 1
            FROM "Note"
            WHERE "EstimateID" = %s
              AND "Category" = 'RFQ'
              AND "NoteText" ILIKE 'All RFQs resolved reminder sent%%'
            LIMIT 1
        """, (estimate_id,))
        return cur.fetchone() is not None
    finally:
        conn.close()

def sweep_for_completed_rfq_estimates():
    """Nudges the owner once all RFQs are resolved but the estimate still has not been sent."""
    from neon_ai.gateway import send_to_user

    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT
                e."EstimateID",
                e."Status",
                e."Description",
                e."CreatedDate",
                c."CustomerName",
                c."Email",
                s."SiteName",
                COUNT(pr."PriceRequestID") AS "RFQCount"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            JOIN "PriceRequest" pr ON pr."EstimateID" = e."EstimateID"
            GROUP BY
                e."EstimateID",
                e."Status",
                e."Description",
                e."CreatedDate",
                c."CustomerName",
                c."Email",
                s."SiteName"
            HAVING COUNT(pr."PriceRequestID") > 0
               AND e."CreatedDate" >= %s
               AND BOOL_AND(COALESCE(pr."Status", '') IN ('Quote Received', 'Declined'))
               AND UPPER(TRIM(COALESCE(e."Status", 'DRAFT'))) NOT IN ('SENT', 'ACCEPTED', 'REJECTED')
            ORDER BY e."EstimateID" ASC
        """, (AUTOMATION_MIN_DATE,))
        rows = cur.fetchall()
    finally:
        conn.close()

    if not rows:
        return

    for row in rows:
        estimate_id = row["EstimateID"]
        if has_rfq_completion_nudge_been_sent(estimate_id):
            continue

        send_to_user(
            subject=f"Estimate Ready To Finish: Estimate #{estimate_id} - {row['CustomerName']}",
            content=(
                f"All RFQs tied to Estimate #{estimate_id} are now resolved as received quotes or declines, "
                "but the estimate has not been sent yet.\n\n"
                f"Customer: {row['CustomerName']}\n"
                f"Customer email: {row['Email'] or 'None on file'}\n"
                f"Site: {row['SiteName']}\n"
                f"Created: {row['CreatedDate']}\n"
                f"Current status: {row['Status']}\n"
                f"RFQs resolved: {row['RFQCount']}\n"
                f"Scope: {row['Description']}\n\n"
                "Next action: finalize the estimate and send it to the customer."
            )
        )
        log_estimate_action(
            estimate_id,
            f"All RFQs resolved reminder sent to owner for Estimate #{estimate_id}.",
            "RFQ"
        )

def has_estimate_age_reminder_been_sent(estimate_id):
    """Checks whether the owner has already been reminded that an estimate hit day 4 unsent."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT 1
            FROM "Note"
            WHERE "EstimateID" = %s
              AND "Category" = 'Follow Up'
              AND "NoteText" ILIKE '4-day estimate preparation reminder sent%%'
            LIMIT 1
        """, (estimate_id,))
        return cur.fetchone() is not None
    finally:
        conn.close()

def sweep_for_aging_unsent_estimates():
    """Reminds the owner when an estimate has existed for 4+ days without being sent."""
    from neon_ai.gateway import send_to_user

    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT
                e."EstimateID",
                e."Status",
                e."Description",
                e."CreatedDate",
                c."CustomerName",
                c."Email",
                s."SiteName"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."CreatedDate" IS NOT NULL
              AND e."CreatedDate" >= %s
              AND e."CreatedDate" <= CURRENT_DATE - 4
              AND UPPER(TRIM(COALESCE(e."Status", 'DRAFT'))) NOT IN ('SENT', 'ACCEPTED', 'REJECTED')
            ORDER BY e."CreatedDate" ASC, e."EstimateID" ASC
        """, (AUTOMATION_MIN_DATE,))
        rows = cur.fetchall()
    finally:
        conn.close()

    if not rows:
        return

    for row in rows:
        estimate_id = row["EstimateID"]
        if has_estimate_age_reminder_been_sent(estimate_id):
            continue

        send_to_user(
            subject=f"4-Day Estimate Reminder: Estimate #{estimate_id} - {row['CustomerName']}",
            content=(
                f"Estimate #{estimate_id} has been open for at least 4 days and still has not been sent.\n\n"
                f"Customer: {row['CustomerName']}\n"
                f"Customer email: {row['Email'] or 'None on file'}\n"
                f"Site: {row['SiteName']}\n"
                f"Created: {row['CreatedDate']}\n"
                f"Current status: {row['Status']}\n"
                f"Scope: {row['Description']}\n\n"
                "Target turnaround is 4 days from request to estimate delivery. Please finish and send this estimate."
            )
        )
        log_estimate_action(
            estimate_id,
            f"4-day estimate preparation reminder sent to owner for Estimate #{estimate_id}.",
            "Follow Up"
        )

def sweep_for_deposit_invoice_reminders():
    """Reminds the owner when an accepted job still has no exported invoice after 24 hours."""
    from neon_ai.gateway import send_to_user

    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT
                e."EstimateID",
                c."CustomerName",
                c."Email",
                s."SiteName",
                MAX(n."Timestamp") AS "AcceptedAt"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            JOIN "Note" n ON n."EstimateID" = e."EstimateID"
            WHERE e."Status" = 'Accepted'
              AND COALESCE(e."SubmitDate", e."CreatedDate") >= %s
              AND n."Category" = 'Customer Response'
              AND n."NoteText" ILIKE %s
            GROUP BY e."EstimateID", c."CustomerName", c."Email", s."SiteName"
            HAVING MAX(n."Timestamp") <= NOW() - INTERVAL '24 hours'
            ORDER BY MAX(n."Timestamp") ASC
        """, (AUTOMATION_MIN_DATE, "%classified as Accepted%"))
        rows = cur.fetchall()
    finally:
        conn.close()

    if not rows:
        print("Sweeper: No deposit invoice reminders needed.")
        return

    for row in rows:
        estimate_id = row["EstimateID"]
        if has_exported_invoice_for_estimate(estimate_id):
            continue
        if has_deposit_invoice_reminder_been_sent(estimate_id):
            continue

        send_to_user(
            subject=f"Deposit Invoice Needed: Estimate #{estimate_id} - {row['CustomerName']}",
            content=(
                f"Estimate #{estimate_id} for {row['CustomerName']} at {row['SiteName']} was accepted more than 24 hours ago, "
                "and no customer invoice has been exported yet.\n\n"
                f"Accepted at: {row['AcceptedAt']}\n"
                f"Customer email: {row['Email']}\n\n"
                "Next action: create and send the deposit invoice for this new job."
            )
        )
        log_estimate_action(
            estimate_id,
            f"24-hour deposit invoice reminder sent to owner. Accepted at {row['AcceptedAt']}.",
            "Billing"
        )

def find_estimate_by_id_and_customer_email(estimate_id, customer_email, statuses=None):
    """Finds a specific estimate by id/email/status combination."""
    statuses = statuses or ["Sent"]
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT e."EstimateID", e."Status", e."Description", s."SiteName", c."CustomerName", c."Email"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."EstimateID" = %s
              AND LOWER(TRIM(c."Email")) = LOWER(TRIM(%s))
              AND e."Status" = ANY(%s)
            LIMIT 1
        """, (estimate_id, customer_email, statuses))
        return cur.fetchone()
    finally:
        conn.close()

def classify_estimate_response(email_subject, email_body):
    """Uses phrase matching to decide whether a reply is an estimate acceptance or rejection."""
    text = f"{email_subject}\n{email_body}".lower()

    rejection_patterns = [
        r"\bnot proceed\b", r"\bwon't proceed\b", r"\bwill not proceed\b",
        r"\bdecline\b", r"\breject\b", r"\btoo expensive\b", r"\btoo much\b",
        r"\bhold off\b", r"\bnot moving forward\b", r"\bpass on this\b"
    ]
    acceptance_patterns = [
        r"\bi accept\b", r"\bwe accept\b", r"\baccepted\b", r"\bapprove\b",
        r"\bapproved\b", r"\bgo ahead\b", r"\bplease proceed\b", r"\bmove forward\b",
        r"\blet's do it\b", r"\blooks good\b", r"\bsounds good\b",
        r"\bok to start\b", r"\bokay to start\b", r"\bsigned\b", r"\bconfirmed\b"
    ]

    if any(re.search(pattern, text) for pattern in rejection_patterns):
        return "Rejected"
    if any(re.search(pattern, text) for pattern in acceptance_patterns):
        return "Accepted"
    return None

def save_post_acceptance_customer_context(customer_email, email_subject, email_body):
    """Stores post-acceptance customer replies in the estimate notes so the PM has scheduling context."""
    accepted_estimates = get_estimates_for_customer_by_status(customer_email, ["Accepted"])
    if not accepted_estimates:
        return None

    estimate_id = extract_estimate_id_from_text(email_subject, email_body)
    if estimate_id:
        estimate = find_estimate_by_id_and_customer_email(estimate_id, customer_email, ["Accepted"])
    elif len(accepted_estimates) == 1:
        estimate = accepted_estimates[0]
    else:
        estimate = accepted_estimates[0]

    if not estimate:
        return None

    note_text = (
        "Post-acceptance customer reply received for scheduling/context.\n"
        f"Subject: {email_subject}\n\n"
        f"{email_body}"
    )
    log_estimate_action(
        estimate["EstimateID"],
        note_text,
        "Customer Scheduling"
    )
    return {
        "estimate_id": estimate["EstimateID"],
        "status": "CustomerSchedulingContext"
    }

def process_estimate_response(customer_email, email_subject, email_body, message_id=None):
    """Processes a customer reply to a previously sent estimate."""
    sent_estimates = get_sent_estimates_for_customer(customer_email)
    if not sent_estimates:
        return save_post_acceptance_customer_context(customer_email, email_subject, email_body)

    estimate_id = extract_estimate_id_from_text(email_subject, email_body)
    if not estimate_id:
        from neon_ai.gateway import send_to_user
        estimate_list = "\n".join(
            f"- Estimate #{row['EstimateID']} for {row['SiteName']}"
            for row in sent_estimates[:5]
        )
        send_to_user(
            subject="Please confirm your estimate number",
            content=(
                f"Hello {sent_estimates[0]['CustomerName']},\n\n"
                f"Thank you for your reply. To make your approval official, please reply again and include the estimate number in the message, "
                f"for example: 'I approve Estimate #{sent_estimates[0]['EstimateID']}'.\n\n"
                f"Our current sent estimates on file are:\n{estimate_list}\n\n"
                f"Best regards,\nArgon Electrical"
            ),
            recipient=customer_email,
            cc_recipients=[os.getenv("PERSONAL_EMAIL")]
        )
        send_to_user(
            subject=f"Estimate Reply Missing Number: {sent_estimates[0]['CustomerName']}",
            content=(
                f"A customer replied without identifying the estimate number.\n\n"
                f"Customer: {sent_estimates[0]['CustomerName']}\n"
                f"Email: {customer_email}\n"
                f"Subject: {email_subject}\n\n"
                f"Sent estimates on file:\n{estimate_list}\n\n"
                f"Original message:\n{email_body}"
            )
        )
        return {"status": "EstimateNumberNeeded"}

    estimate = find_estimate_by_id_and_customer_email(estimate_id, customer_email, ["Sent"])
    if not estimate:
        from neon_ai.gateway import send_to_user
        send_to_user(
            subject="Estimate number could not be confirmed",
            content=(
                "Hello,\n\n"
                f"Thank you for your reply. We could not match Estimate #{estimate_id} to an active sent estimate for this email address.\n\n"
                "Please reply again with the correct estimate number exactly as shown on your estimate document.\n\n"
                "Best regards,\nArgon Electrical"
            ),
            recipient=customer_email,
            cc_recipients=[os.getenv("PERSONAL_EMAIL")]
        )
        send_to_user(
            subject=f"Estimate Number Mismatch: {customer_email}",
            content=(
                f"A customer replied with Estimate #{estimate_id}, but it did not match an active sent estimate for their email.\n\n"
                f"Subject: {email_subject}\n\n"
                f"Message:\n{email_body}"
            )
        )
        return {"status": "EstimateNumberMismatch", "estimate_id": estimate_id}

    response_type = classify_estimate_response(email_subject, email_body)
    if not response_type:
        return None

    from neon_ai.gateway import send_to_user

    archive_path = None
    before_snapshot = get_estimate_verification_snapshot(estimate["EstimateID"])
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            'UPDATE "Estimate" SET "Status" = %s WHERE "EstimateID" = %s',
            (response_type, estimate["EstimateID"])
        )
        conn.commit()
    finally:
        conn.close()

    if response_type == "Accepted":
        try:
            archive_path = archive_estimate_response_pdf(
                estimate["EstimateID"],
                customer_email,
                email_subject,
                email_body,
                message_id=message_id,
                response_type=response_type
            )
        except Exception as exc:
            archive_path = None
            log_estimate_action(
                estimate["EstimateID"],
                f"Acceptance email PDF archive failed: {exc}",
                "System"
            )

    log_estimate_action(
        estimate["EstimateID"],
        f"Customer email response classified as {response_type}. Subject: {email_subject}",
        "Customer Response"
    )
    after_snapshot = get_estimate_verification_snapshot(estimate["EstimateID"])
    log_estimate_verification(
        estimate["EstimateID"],
        f"customer estimate response -> {response_type}",
        before_snapshot,
        after_snapshot,
        expected_values={"EstimateStatus": response_type},
        expected_changed_fields=["EstimateStatus"],
    )

    if response_type == "Accepted":
        record_estimate_workflow_handoff(
            estimate["EstimateID"],
            "ConvertEstimateToWorkOrder",
            proof_type="Email",
            proof_reference=archive_path or message_id or email_subject[:255]
        )
        if archive_path:
            log_estimate_action(
                estimate["EstimateID"],
                f"Acceptance email archived to PDF at {archive_path}.",
                "Customer Response"
            )
        send_to_user(
            subject=f"Thank you - estimate accepted for {estimate['SiteName']}",
            content=(
                f"Hello {estimate['CustomerName']},\n\n"
                f"Thank you for approving Estimate #{estimate['EstimateID']} for {estimate['SiteName']}.\n\n"
                "We appreciate the opportunity to do the work for you.\n\n"
                "Please let us know when you would ideally like the work to begin, and we will coordinate the next steps from there.\n\n"
                f"Best regards,\nArgon Electrical"
            ),
            recipient=customer_email,
            cc_recipients=[os.getenv("PERSONAL_EMAIL")]
        )
        send_to_user(
            subject=f"Estimate Accepted: {estimate['CustomerName']} - {estimate['SiteName']}",
            content=(
                f"Customer acceptance received for Estimate #{estimate['EstimateID']}.\n\n"
                f"Customer: {estimate['CustomerName']}\n"
                f"Site: {estimate['SiteName']}\n"
                f"Email: {customer_email}\n"
                f"Archive: {archive_path}\n"
                f"Subject: {email_subject}\n\n"
                f"Message:\n{email_body}\n\n"
                "Next action: create and send the deposit invoice for this new job, then continue work-order intake."
            )
        )
    else:
        send_to_user(
            subject=f"Estimate Response: {estimate['CustomerName']} - {response_type}",
            content=(
                f"Estimate #{estimate['EstimateID']} was classified as {response_type} from customer reply.\n\n"
                f"Customer: {estimate['CustomerName']}\n"
                f"Site: {estimate['SiteName']}\n"
                f"Email: {customer_email}\n"
                f"Subject: {email_subject}\n\n"
                f"Message:\n{email_body}"
            )
        )

    return {"estimate_id": estimate["EstimateID"], "status": response_type}

# ==========================================
# 2. INTAKE & LEAD GENERATION
# ==========================================

def create_new_lead(customer_name, email, phone, inquiry_text, address=None, scope=None):
    """Intelligent Upsert: Creates a new lead OR updates an existing open Draft Estimate."""
    from neon_ai.database.connection import get_connection
    from psycopg2.extras import RealDictCursor
    
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor) # Using RealDictCursor to fetch by column name
    try:
        # 1. FIND OR CREATE CUSTOMER
        clean_email = email.split('<')[-1].strip('>') if '<' in email else email.strip()
        cur.execute('SELECT "CustomerID" FROM "Customer" WHERE "Email" ILIKE %s', (f"%{clean_email}%",))
        existing_customer = cur.fetchone()

        if existing_customer:
            customer_id = existing_customer['CustomerID']
            if phone and str(phone).lower() not in ["pending", "null", "none"]:
                 cur.execute('UPDATE "Customer" SET "Phone" = %s WHERE "CustomerID" = %s AND ("Phone" IS NULL OR "Phone" = \'Pending\')', (phone, customer_id))
        else:
            cur.execute('INSERT INTO "Customer" ("CustomerName", "Email", "Phone") VALUES (%s, %s, %s) RETURNING "CustomerID"', 
                        (customer_name, clean_email, phone))
            customer_id = cur.fetchone()['CustomerID']

        # 2. CHECK FOR EXISTING OPEN DRAFT
        cur.execute('''
            SELECT e."EstimateID", s."SiteID", s."SiteName"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            WHERE s."CustomerID" = %s AND e."Status" = 'DRAFT'
            ORDER BY e."EstimateID" DESC LIMIT 1
        ''', (customer_id,))
        open_draft = cur.fetchone()

        actual_scope = scope if scope and str(scope).lower() not in ["null", "none"] else inquiry_text[:200]
        before_snapshot = get_estimate_verification_snapshot(open_draft['EstimateID']) if open_draft else None

        if open_draft:
            # 3A. UPDATE THE EXISTING DRAFT
            estimate_id = open_draft['EstimateID']
            site_id = open_draft['SiteID']
            current_site_name = open_draft['SiteName']
            
            if address and str(address).lower() not in ["null", "none"] and "Pending Site" in current_site_name:
                cur.execute('UPDATE "Site" SET "SiteName" = %s WHERE "SiteID" = %s', (address, site_id))
            
            cur.execute('''
                UPDATE "Estimate" 
                SET "Description" = CONCAT("Description", ' | Update: ', %s) 
                WHERE "EstimateID" = %s
            ''', (actual_scope, estimate_id))
            
            cur.execute('INSERT INTO "Note" ("EstimateID", "NoteText", "Category") VALUES (%s, %s, \'Lead Update\')', (estimate_id, inquiry_text))
            action_status = "UPDATED"

        else:
            # 3B. CREATE NEW SITE AND ESTIMATE
            valid_address = address if address and str(address).lower() not in ["null", "none"] else f"Pending Site - {customer_name}"
            sync_table_sequence(cur, "Site", "SiteID")
            cur.execute('INSERT INTO "Site" ("CustomerID", "SiteName") VALUES (%s, %s) RETURNING "SiteID"', (customer_id, valid_address))
            site_id = cur.fetchone()['SiteID']

            ensure_estimate_sequence(cur)
            cur.execute('INSERT INTO "Estimate" ("SiteID", "Description", "Status") VALUES (%s, %s, \'DRAFT\') RETURNING "EstimateID"', (site_id, actual_scope))
            estimate_id = cur.fetchone()['EstimateID']

            cur.execute('INSERT INTO "Note" ("EstimateID", "NoteText", "Category") VALUES (%s, %s, \'Lead Inquiry\')', (estimate_id, inquiry_text))
            action_status = "CREATED"

        # Final audit: verify the customer intake record is actually complete.
        cur.execute('''
            SELECT
                c."CustomerID",
                c."CustomerName",
                c."Email",
                c."Phone",
                c."Address",
                c."StreetName",
                c."CityName",
                c."PostalCode",
                s."SiteName"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."EstimateID" = %s
        ''', (estimate_id,))
        final_record = cur.fetchone()

        missing_fields = []
        if not final_record['CustomerName'] or str(final_record['CustomerName']).lower() in ["unknown", "null", "none"]:
            missing_fields.append("CustomerName")
        if not final_record['Email'] or str(final_record['Email']).lower() in ["pending", "null", "none"]:
            missing_fields.append("CustomerEmail")
        if not final_record['Phone'] or str(final_record['Phone']).lower() in ["pending", "null", "none"]:
            missing_fields.append("CustomerPhone")
        if not final_record['Address'] or str(final_record['Address']).lower() in ["pending", "null", "none"]:
            missing_fields.append("AddressNumber")
        if not final_record['StreetName'] or str(final_record['StreetName']).lower() in ["pending", "null", "none"]:
            missing_fields.append("StreetName")
        if not final_record['CityName'] or str(final_record['CityName']).lower() in ["pending", "null", "none"]:
            missing_fields.append("CityName")
        if not final_record['PostalCode'] or str(final_record['PostalCode']).lower() in ["pending", "null", "none"]:
            missing_fields.append("PostalCode")
        if not final_record['SiteName'] or "Pending Site" in final_record['SiteName']:
            missing_fields.append("SiteAddress")

        conn.commit()
        after_snapshot = get_estimate_verification_snapshot(estimate_id)
        log_estimate_verification(
            estimate_id,
            f"create_new_lead {action_status}",
            before_snapshot,
            after_snapshot,
            expected_values={"EstimateStatus": "DRAFT"},
        )
        
        # Return the 'missing' array back to the Gateway!
        return {
            "status": action_status,
            "estimate_id": estimate_id,
            "customer_id": final_record['CustomerID'],
            "name": final_record['CustomerName'],
            "missing": missing_fields
        }
        
    except Exception as e:
        conn.rollback()
        print(f"DATABASE ERROR: {e}")
        return {"status": "ERROR", "message": str(e)}
    finally:
        conn.close()

def notify_milton_of_ready_lead(estimate_id: int):
    """Generates the folder, saves the Word doc (Lead Sheet), alerts Milton, and logs it."""
    from neon_ai.gateway import send_to_user 
    from neon_ai.database.connection import get_connection
    from psycopg2.extras import RealDictCursor
    from docx import Document
    import os

    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        # 1. Fetch Core Lead Info (Added Phone and Email)
        cur.execute('''
            SELECT c."CustomerName", c."Email", c."Phone", s."SiteName", e."Description"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."EstimateID" = %s
        ''', (estimate_id,))
        job_data = cur.fetchone()

        if not job_data:
            print(f"⚠️ Error: Could not find Estimate #{estimate_id} in database.")
            return "ERROR: Missing DB Record"

        # 2. Fetch ALL Notes for this Estimate (Ordered by oldest to newest)
        cur.execute('''
            SELECT "NoteText", "Category", "Timestamp"
            FROM "Note"
            WHERE "EstimateID" = %s
            ORDER BY "Timestamp" ASC
        ''', (estimate_id,))
        notes_data = cur.fetchall()

    finally:
        conn.close()

    name = job_data['CustomerName']
    site_address = job_data['SiteName']
    scope = job_data['Description']
    phone = job_data['Phone'] or "Not Provided"
    email = job_data['Email'] or "Not Provided"

    # THE GOLDEN RULE CALL
    folder_path, docx_filename = get_project_file_paths(name, site_address, estimate_id)

    # 3. Build the Lead Sheet Document
    doc = Document()
    doc.add_heading(f"Argon Lead Sheet: {name}", 0)
    
    doc.add_heading("Customer Information", level=1)
    doc.add_paragraph(f"Name: {name}\nPhone: {phone}\nEmail: {email}")

    doc.add_heading("Site Information", level=1)
    doc.add_paragraph(f"Address: {site_address}\nEstimate Ref: #{estimate_id}")

    doc.add_heading("Scope of Work", level=1)
    doc.add_paragraph(scope)

    doc.add_heading("Audit Trail & Notes", level=1)
    if notes_data:
        for note in notes_data:
            # Format: [Date] Category: Text
            date_str = note['Timestamp'].strftime("%Y-%m-%d %H:%M") if note['Timestamp'] else "Unknown Date"
            doc.add_paragraph(f"[{date_str}] {note['Category']}: {note['NoteText']}")
    else:
        doc.add_paragraph("No notes recorded yet.")

    doc.save(docx_filename)

    # 4. Send the Email
    body = f"🚀 LEAD SHEET READY: {name}\n\nAddress: {site_address}\nScope: {scope}\n\nSaved to: {folder_path}"
    send_to_user(f"Action Required: Site Walk for {name}", body, attachment_path=docx_filename)
    
    # 5. --- THE PEN: Write it down! ---
    log_estimate_action(estimate_id, f"Lead complete. Lead Sheet generated and Milton notified at {folder_path}.")
    
    return f"SUCCESS: Archived to {folder_path}"

def send_smart_discovery(estimate_id: int, customer_name: str, email_address: str, missing_info: str) -> str:
    """Stages and SENDS a 'Smart Discovery' email, then logs it to the database."""
    from neon_ai.gateway import send_to_user # Avoid circular import
    
    try:
        subject = f"Information needed for your Argon Electrical Estimate - {customer_name}"
        draft_body = (
            f"Hello {customer_name},\n\n"
            f"Thank you for reaching out to Argon Electrical! To get you an accurate estimate, "
            f"we just need a few more details regarding your project:\n\n"
            f"- {missing_info}\n\n"
            f"Please reply to this email with the information (and any photos of the area if applicable), "
            f"and we will get a quote started for you right away.\n\n"
            f"Best regards,\nArgon Electrical"
        )
        
        # Actually send the email
        success = send_to_user(subject, draft_body, recipient=email_address)
        
        if success:
            print(f"\n✅ EMAIL SENT: Smart Discovery sent to {email_address}")
            # --- THE PEN: Write it down! ---
            log_estimate_action(estimate_id, f"Sent Smart Discovery asking for: {missing_info}")
            return f"Success: Discovery email sent."
        else:
            return f"Error: SMTP failed to send to {email_address}."
            
    except Exception as e:
        print(f"Discovery Tool Error: {e}")
        return f"Error: Could not send discovery email. {e}"


def send_customer_info_request(estimate_id: int, customer_name: str, email_address: str, missing_fields):
    """Politely request any customer-profile fields still missing after inbox backfill."""
    from neon_ai.gateway import send_to_user

    friendly_fields = format_customer_missing_fields(missing_fields)
    bullet_lines = "\n".join(f"- {field}" for field in friendly_fields)
    name = customer_name or "there"

    subject = f"Quick customer info check for your Argon estimate"
    body = (
        f"Hello {name},\n\n"
        "We are getting your customer record organized so we can keep your estimate and project details accurate.\n\n"
        "Could you please reply with the following missing information:\n"
        f"{bullet_lines}\n\n"
        "If this information was already sent, feel free to just reply with the details again in one message and we will update the file.\n\n"
        "Thank you,\nArgon Electrical"
    )

    if send_to_user(subject, body, recipient=email_address):
        log_estimate_action(
            estimate_id,
            f"Customer info request sent asking for: {', '.join(missing_fields)}",
            "Customer Intake",
        )
        return True
    return False

# ==========================================
# 3. OUTBOX SWEEPER (FOR LOCKED JOBS)
# ==========================================

def _deprecated_legacy_sweep_locked_estimates_to_drafts():
    """Deprecated legacy draft-based sweeper. Do not call."""
    from neon_ai.gateway import _deprecated_save_to_drafts_folder, send_to_user
    from neon_ai.automation.local_brain import ArgonLocalAI 
    
    conn = get_connection(); cur = conn.cursor(); ai = ArgonLocalAI()

    try:
        cur.execute("""
            SELECT e."EstimateID", e."Description", e."TotalAmount", 
                   c."CustomerName", c."Email", s."SiteName"
            FROM "Estimate" e
            JOIN "Site" s ON e."SiteID" = s."SiteID"
            JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
            WHERE e."Status" = 'Locked'
        """)
        locked_estimates = cur.fetchall()

        if not locked_estimates:
            print("💤 Sweeper: No 'Locked' estimates found."); return 

        for est in locked_estimates:
            folder_path, file_path = get_project_file_paths(est['CustomerName'], est['SiteName'], est['EstimateID'])

            if not os.path.exists(file_path):
                print(f"⚠️ Warning: File not found at {file_path}. Skipping."); continue

            email_body = ai.draft_proposal_email(est['CustomerName'], est['SiteName'], est['Description'], float(est['TotalAmount']))
            
            subject = f"Estimate from Argon Electrical: {est['SiteName']}"
            if _deprecated_save_to_drafts_folder(subject, email_body, est['Email'], attachment_path=file_path):
                send_to_user(f"Ready to Send: {est['CustomerName']} Quote", f"Draft is in your Outbox. File: {file_path}")
                cur.execute('UPDATE "Estimate" SET "Status" = \'Pending Review\' WHERE "EstimateID" = %s', (est['EstimateID'],))
                conn.commit()
                # Log the Sweeper Action!
                log_estimate_action(est['EstimateID'], "Sweeper moved locked estimate to Gmail Drafts.", "System")
    finally: conn.close()

# Override the legacy draft-saver with the live sender workflow.
def _legacy_estimate_auto_send_enabled() -> bool:
    return str(os.getenv(LEGACY_ESTIMATE_AUTO_SEND_ENV, "0")).strip().lower() in {"1", "true", "yes", "on"}


def _run_legacy_locked_estimate_auto_send(cur) -> None:
    """Legacy customer-send behavior retained only behind an explicit env gate."""
    from neon_ai.gateway import send_to_user

    cur.execute("""
        SELECT e."EstimateID", e."Description", e."TotalAmount",
               c."CustomerName", c."Email", s."SiteName"
        FROM "Estimate" e
        JOIN "Site" s ON e."SiteID" = s."SiteID"
        JOIN "Customer" c ON s."CustomerID" = c."CustomerID"
        WHERE e."Status" = 'Locked'
          AND e."CreatedDate" >= %s
    """, (AUTOMATION_MIN_DATE,))
    locked_estimates = cur.fetchall()

    if not locked_estimates:
        print("Sweeper: No 'Locked' estimates found.")
        return

    print(
        f"Sweeper: {LEGACY_ESTIMATE_AUTO_SEND_ENV}=1 detected. "
        "Running legacy estimate auto-send behavior."
    )
    for est in locked_estimates:
        file_path = build_client_estimate_doc(est['EstimateID'])
        total_amount = get_client_estimate_total(est['EstimateID'])
        email_body = (
            f"Hello {est['CustomerName']},\n\n"
            f"Please find attached Estimate #{est['EstimateID']} for {est['SiteName']}.\n\n"
            f"Project scope:\n{est['Description']}\n\n"
            f"Estimated total: ${total_amount:,.2f}\n\n"
            f"If you'd like to move forward, please reply to this email and include the exact wording 'I approve Estimate #{est['EstimateID']}' so we can confirm it officially.\n\n"
            f"Best regards,\nArgon Electrical"
        )
        subject = f"Estimate #{est['EstimateID']} from Argon Electrical: {est['SiteName']}"

        if send_to_user(
            subject,
            email_body,
            recipient=est['Email'],
            attachment_path=file_path,
            cc_recipients=[os.getenv("PERSONAL_EMAIL")]
        ):
            cur.execute(
                'UPDATE "Estimate" SET "Status" = \'Sent\', "SubmitDate" = CURRENT_DATE WHERE "EstimateID" = %s',
                (est['EstimateID'],)
            )
            log_estimate_action(
                est['EstimateID'],
                f"Locked estimate auto-sent to customer and CC'd to owner. SubmitDate stamped. Attachment: {file_path}",
                "System"
            )


def sweep_for_locked_estimates():
    """Fence legacy estimate auto-send so locked estimates require explicit Final Doc View send."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        if _legacy_estimate_auto_send_enabled():
            _run_legacy_locked_estimate_auto_send(cur)
            conn.commit()
            return

        cur.execute("""
            SELECT COUNT(*) AS "LockedCount"
            FROM "Estimate" e
            WHERE e."Status" = 'Locked'
              AND e."CreatedDate" >= %s
        """, (AUTOMATION_MIN_DATE,))
        row = cur.fetchone() or {}
        locked_count = int(row.get("LockedCount") or 0)
        if locked_count <= 0:
            print("Sweeper: No 'Locked' estimates found.")
            return

        print(
            "Sweeper: Legacy estimate auto-send is disabled. "
            "Use Final Doc View Send to Customer. "
            f"Locked estimates pending manual send: {locked_count}."
        )
    finally:
        conn.close()

# ==========================================
# 4. AI AGENT TOOLS (WORK ORDERS)
# ==========================================

def get_project_memory(workorder_id: int) -> str:
    """Retrieves the core status and details of a specific project/work order."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute('''
            SELECT "WorkOrderID", "JobStatus", "Description", "BillingType", "CustomerPO"
            FROM public."WorkOrder"
            WHERE "WorkOrderID" = %s
        ''', (workorder_id,))
        
        project = cur.fetchone()
        if not project: return f"Error: No project found with ID {workorder_id}."

        return (
            f"Project #{project['WorkOrderID']} Details:\n"
            f"- Current Status: {project['JobStatus'] or 'Unknown'}\n"
            f"- Billing Type: {project['BillingType'] or 'Unknown'}\n"
            f"- Customer PO: {project['CustomerPO'] or 'Not Provided'}\n"
            f"- Scope: {project['Description'] or 'No description on file'}\n"
        )
    except Exception as e:
        return f"Database Error: Could not retrieve memory. {e}"
    finally: conn.close()

def add_note_to_active_job(workorder_id: int, note_text: str) -> str:
    """Adds a note to the source estimate behind an active work order."""
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute('''
            SELECT "SourceEstimateID"
            FROM public."WorkOrder"
            WHERE "WorkOrderID" = %s
        ''', (workorder_id,))
        row = cur.fetchone()
        if not row or not row.get("SourceEstimateID"):
            return f"Database Error: Work Order #{workorder_id} is not linked to a source estimate."

        cur.execute('''
            INSERT INTO public."Note" ("EstimateID", "NoteText", "Category")
            VALUES (%s, %s, %s)
        ''', (row["SourceEstimateID"], note_text, 'Work Order'))
        conn.commit()
        return f"Success: Note successfully added to Work Order #{workorder_id}."
    except Exception as e:
        return f"Database Error: Could not add note. {e}"
    finally: conn.close()

# 5. THE INBOX SWEEPER (FOR NEW LEADS)
# ==========================================

def mark_estimate_ready(estimate_id: int):
    """Lightning-fast update so the Gateway can hand off the heavy lifting."""
    from neon_ai.database.connection import get_connection
    before_snapshot = get_estimate_verification_snapshot(estimate_id)
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute('UPDATE "Estimate" SET "Status" = \'READY\' WHERE "EstimateID" = %s', (estimate_id,))
        conn.commit()
        after_snapshot = get_estimate_verification_snapshot(estimate_id)
        log_estimate_verification(
            estimate_id,
            "mark_estimate_ready",
            before_snapshot,
            after_snapshot,
            expected_values={"EstimateStatus": "READY"},
            expected_changed_fields=["EstimateStatus"],
        )
    except Exception as e:
        print(f"Database Error marking ready: {e}")
    finally:
        conn.close()

def sweep_for_ready_leads():
    """Hunts for 'READY' estimates, generates Lead Sheets, and resets to DRAFT."""
    from neon_ai.database.connection import get_connection
    from psycopg2.extras import RealDictCursor
    # Note: We already have notify_milton_of_ready_lead in this file, so we can just call it!
    
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            'SELECT "EstimateID" FROM "Estimate" WHERE "Status" = \'READY\' AND COALESCE("CreatedDate", CURRENT_DATE) >= %s',
            (AUTOMATION_MIN_DATE,)
        )
        ready_estimates = cur.fetchall()

        if not ready_estimates:
            return 

        for est in ready_estimates:
            est_id = est['EstimateID']
            print(f"\n🧹 [SWEEPER] Found Ready Lead #{est_id}. Generating Lead Sheet...")
            
            # Do the heavy lifting (Build doc, send email)
            notify_milton_of_ready_lead(est_id)
            
            # Flip it back to DRAFT so it doesn't get swept again, and is ready for the UI
            cur.execute('UPDATE "Estimate" SET "Status" = \'DRAFT\' WHERE "EstimateID" = %s', (est_id,))
            conn.commit()
            print(f"✅ [SWEEPER] Lead #{est_id} safely moved back to DRAFT state.")
            
    except Exception as e:
        print(f"Sweeper Error: {e}")
    finally:
        conn.close()
