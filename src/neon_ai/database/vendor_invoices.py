import datetime
import json
import os
import re
import time
from collections import Counter
from decimal import Decimal

from docx import Document
from psycopg2.extras import Json, RealDictCursor

from neon_ai.automation.runtime_flags import require_legacy_automation_runtime
from neon_ai.database.connection import get_connection
from neon_ai.database.folders import get_target_folder
from neon_ai.database.purchases import (
    AUTOMATION_MIN_DATE,
    _extract_numeric_values,
    _match_po_item_to_line,
    ensure_purchase_schema,
    get_po_export_data,
    get_purchase_order_receipt_choices,
    get_purchase_order_receipt_quantities,
    update_purchase_order_receiving_status,
)
from neon_ai.database.rfq import extract_text_from_attachment, ocr_pdf_attachment

VENDOR_INVOICE_MEMORY_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "vendor_invoice_memory.json")
_VENDOR_INVOICE_SCHEMA_READY = False


def _perf_log(area: str, name: str, started_at: float) -> None:
    elapsed_ms = (time.perf_counter() - started_at) * 1000.0
    print(f"[PERF] area={area} name={name} elapsed_ms={elapsed_ms:.2f}")


def ensure_vendor_invoice_schema():
    global _VENDOR_INVOICE_SCHEMA_READY
    if _VENDOR_INVOICE_SCHEMA_READY:
        return
    ensure_purchase_schema()
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute('ALTER TABLE "VendorInvoice" ADD COLUMN IF NOT EXISTS "VendorInvoicePaidDate" date')
        cur.execute('ALTER TABLE "VendorInvoice" ADD COLUMN IF NOT EXISTS "VendorInvoiceMemory" jsonb')
        conn.commit()
        _VENDOR_INVOICE_SCHEMA_READY = True
    finally:
        conn.close()


def _load_legacy_vendor_invoice_memory():
    if not os.path.exists(VENDOR_INVOICE_MEMORY_PATH):
        return []
    try:
        with open(VENDOR_INVOICE_MEMORY_PATH, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        return data if isinstance(data, list) else []
    except Exception:
        return []


def _json_db_ready(value):
    if isinstance(value, dict):
        return {str(key): _json_db_ready(val) for key, val in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_db_ready(item) for item in value]
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def _decode_vendor_invoice_memory(value):
    if isinstance(value, dict):
        return value
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
            return parsed if isinstance(parsed, dict) else {}
        except Exception:
            return {}
    return {}


def _read_vendor_invoice_memory_from_db(invoice_id: int):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            'SELECT "VendorInvoiceMemory" FROM "VendorInvoice" WHERE "VendorInvoiceID" = %s',
            (invoice_id,),
        )
        row = cur.fetchone() or {}
        return _decode_vendor_invoice_memory(row.get("VendorInvoiceMemory"))
    finally:
        conn.close()


def _write_vendor_invoice_memory_to_db(invoice_id: int, payload: dict):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            UPDATE "VendorInvoice"
            SET "VendorInvoiceMemory" = %s
            WHERE "VendorInvoiceID" = %s
            ''',
            (Json(_json_db_ready(payload or {})), invoice_id),
        )
        conn.commit()
    finally:
        conn.close()


def load_vendor_invoice_memory():
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT "VendorInvoiceID", "VendorInvoiceMemory"
            FROM "VendorInvoice"
            WHERE "VendorInvoiceMemory" IS NOT NULL
            ORDER BY "VendorInvoiceID" ASC
            '''
        )
        rows = cur.fetchall()
    finally:
        conn.close()

    entries = []
    for row in rows:
        payload = _decode_vendor_invoice_memory(row.get("VendorInvoiceMemory"))
        if payload:
            payload.setdefault("vendor_invoice_id", row.get("VendorInvoiceID"))
            entries.append(payload)
    return entries or _load_legacy_vendor_invoice_memory()


def save_vendor_invoice_memory(entries):
    for entry in entries or []:
        invoice_id = int(entry.get("vendor_invoice_id") or 0)
        if invoice_id:
            _write_vendor_invoice_memory_to_db(invoice_id, entry)


def upsert_vendor_invoice_memory(invoice_id: int, updates: dict):
    entry = dict(get_vendor_invoice_memory(invoice_id) or {})
    entry["vendor_invoice_id"] = int(invoice_id)
    entry.update(_json_db_ready(updates or {}))
    _write_vendor_invoice_memory_to_db(invoice_id, entry)
    return entry


def get_vendor_invoice_memory(invoice_id: int):
    payload = _read_vendor_invoice_memory_from_db(invoice_id)
    if payload:
        payload.setdefault("vendor_invoice_id", int(invoice_id))
        return payload

    legacy = next((row for row in _load_legacy_vendor_invoice_memory() if int(row.get("vendor_invoice_id") or 0) == int(invoice_id)), {})
    if legacy:
        legacy = _json_db_ready(dict(legacy))
        legacy["vendor_invoice_id"] = int(invoice_id)
        _write_vendor_invoice_memory_to_db(invoice_id, legacy)
    return legacy


def _parse_invoice_date(text: str, reference_date=None):
    reference_date = reference_date or datetime.date.today()
    if not text:
        return None

    iso_match = re.search(r'\b(20\d{2})-(\d{1,2})-(\d{1,2})\b', text)
    if iso_match:
        y, m, d = map(int, iso_match.groups())
        return datetime.date(y, m, d)

    slash_match = re.search(r'\b(\d{1,2})[/-](\d{1,2})[/-](20\d{2}|\d{2})\b', text)
    if slash_match:
        m, d, y = slash_match.groups()
        year = int(y)
        year = year + 2000 if year < 100 else year
        return datetime.date(year, int(m), int(d))

    month_match = re.search(r'\b(january|february|march|april|may|june|july|august|september|october|november|december)\s+(\d{1,2})(?:st|nd|rd|th)?(?:,\s*(20\d{2}))?', text, re.IGNORECASE)
    if month_match:
        month_name, day_value, year_value = month_match.groups()
        month = datetime.datetime.strptime(month_name.title(), "%B").month
        year = int(year_value) if year_value else reference_date.year
        return datetime.date(year, month, int(day_value))

    return None


def extract_invoice_header_fields(subject: str, body: str, attachment_text: str = "", reference_date=None):
    combined = "\n".join(part for part in [subject, body, attachment_text] if part)
    invoice_number = None
    amount = None
    invoice_date = _parse_invoice_date(combined, reference_date=reference_date)

    inv_match = re.search(r'\b(?:invoice|inv)\s*#?\s*([A-Za-z0-9._/-]+)\b', combined, re.IGNORECASE)
    if inv_match:
        invoice_number = inv_match.group(1)[:80]

    total_patterns = [
        r'\btotal(?: due)?[:\s$]*([\d,]+\.\d{2})',
        r'\bbalance(?: due)?[:\s$]*([\d,]+\.\d{2})',
        r'\binvoice amount[:\s$]*([\d,]+\.\d{2})',
    ]
    for pattern in total_patterns:
        match = re.search(pattern, combined, re.IGNORECASE)
        if match:
            amount = float(match.group(1).replace(",", ""))
            break

    if amount is None:
        numeric_values = _extract_numeric_values(combined)
        money_like = [val for val in numeric_values if val >= 1]
        if money_like:
            amount = max(money_like)

    return {
        "invoice_number": invoice_number or "",
        "invoice_date": invoice_date.isoformat() if invoice_date else (reference_date.isoformat() if reference_date else datetime.date.today().isoformat()),
        "amount": amount or 0.0,
    }


def _get_po_item_rows(po_id: int):
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT
                poi."POItemID",
                poi."PurchaseOrderID",
                poi."MaterialID",
                poi."CatalogItemID",
                COALESCE(poi."Description", em."Description", m."Description") AS "Description",
                poi."QuantityOrdered",
                COALESCE(poi."QuantityReceived", 0) AS "QuantityReceived",
                COALESCE(
                    NULLIF(poi."UnitPriceAtOrder", 0),
                    NULLIF(em."CommittedUnitCost", 0),
                    NULLIF(em."AwardedUnitCost", 0),
                    NULLIF(em."UnitCost", 0),
                    CASE COALESCE(m."CarryPriceSource", 'Internal')
                        WHEN 'Nedco' THEN NULLIF(m."NedcoPrice", 0)
                        WHEN 'Gescan' THEN NULLIF(m."GescanPrice", 0)
                        WHEN 'Eecol' THEN NULLIF(m."EecolPrice", 0)
                        WHEN 'Guillevin' THEN NULLIF(m."GuillevinPrice", 0)
                        ELSE NULLIF(m."InternalPrice", 0)
                    END,
                    0
                ) AS "POUnitPrice",
                COALESCE(
                    NULLIF(poi."LineTotal", 0),
                    COALESCE(poi."QuantityOrdered", 0) * COALESCE(
                        NULLIF(poi."UnitPriceAtOrder", 0),
                        NULLIF(em."CommittedUnitCost", 0),
                        NULLIF(em."AwardedUnitCost", 0),
                        NULLIF(em."UnitCost", 0),
                        CASE COALESCE(m."CarryPriceSource", 'Internal')
                            WHEN 'Nedco' THEN NULLIF(m."NedcoPrice", 0)
                            WHEN 'Gescan' THEN NULLIF(m."GescanPrice", 0)
                            WHEN 'Eecol' THEN NULLIF(m."EecolPrice", 0)
                            WHEN 'Guillevin' THEN NULLIF(m."GuillevinPrice", 0)
                            ELSE NULLIF(m."InternalPrice", 0)
                        END,
                        0
                    ),
                    0
                ) AS "CurrentLineTotal",
                COALESCE(
                    em."UnitCost",
                    CASE COALESCE(m."CarryPriceSource", 'Internal')
                        WHEN 'Nedco' THEN m."NedcoPrice"
                        WHEN 'Gescan' THEN m."GescanPrice"
                        WHEN 'Eecol' THEN m."EecolPrice"
                        WHEN 'Guillevin' THEN m."GuillevinPrice"
                        ELSE m."InternalPrice"
                    END,
                    0
                ) AS "EstimateUnitPrice",
                (COALESCE(poi."QuantityOrdered", 0) - COALESCE(poi."QuantityReceived", 0)) AS "BackorderedQty"
            FROM "PurchaseOrderItem" poi
            LEFT JOIN "EstimateMaterial" em ON poi."MaterialID" = em."EstimateMaterialID"
            LEFT JOIN "Material" m ON poi."CatalogItemID" = m."ItemID"
            WHERE poi."PurchaseOrderID" = %s
            ORDER BY poi."POItemID" ASC
            ''',
            (po_id,),
        )
        return cur.fetchall()
    finally:
        conn.close()


def parse_vendor_invoice_lines(po_id: int, text: str):
    po_items = _get_po_item_rows(po_id)
    if not po_items:
        return []

    parsed = []
    for raw_line in (text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        item = _match_po_item_to_line(line, po_items)
        if not item:
            continue

        numbers = _extract_numeric_values(line)
        qty = numbers[0] if numbers else float(item.get("QuantityReceived") or item.get("QuantityOrdered") or 0)
        unit_price = numbers[-1] if len(numbers) >= 2 else float(item.get("POUnitPrice") or 0)
        parsed.append(
            {
                "po_item_id": item["POItemID"],
                "description": item.get("Description"),
                "invoice_qty": qty,
                "invoice_unit_price": unit_price,
                "invoice_line_total": qty * unit_price,
                "source_line": line[:255],
            }
        )

    deduped = {}
    for row in parsed:
        deduped[row["po_item_id"]] = row
    return list(deduped.values())


def find_matching_po_for_vendor_invoice(sender_email: str, subject: str = "", body: str = "", attachment_text: str = ""):
    combined = "\n".join(part for part in [subject, body, attachment_text] if part)
    if "invoice" not in combined.lower() and "statement" not in combined.lower() and "bill" not in combined.lower():
        return None

    po_match = re.search(r'\bpo\s*#?\s*(\d+)\b', combined, re.IGNORECASE)
    if po_match:
        return int(po_match.group(1))

    sender = (sender_email or "").strip().lower()
    if not sender:
        return None

    from neon_ai.database.purchases import load_po_dispatch_log

    matches = [row for row in load_po_dispatch_log() if (row.get("recipient_email") or "").strip().lower() == sender]
    if not matches:
        return None
    matches.sort(key=lambda row: row.get("sent_at", ""), reverse=True)
    return matches[0].get("po_id")


def _get_latest_vendor_invoice_for_po(po_id: int, statuses=None):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        sql = '''
            SELECT *
            FROM "VendorInvoice"
            WHERE "PurchaseOrderID" = %s
        '''
        params = [po_id]
        if statuses:
            sql += ' AND "VendorInvoiceStatus" = ANY(%s)'
            params.append(statuses)
        sql += ' ORDER BY "VendorInvoiceID" DESC LIMIT 1'
        cur.execute(sql, params)
        return cur.fetchone()
    finally:
        conn.close()


def _find_vendor_invoice_by_number(po_id: int, invoice_number: str):
    if not invoice_number:
        return None
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT *
            FROM "VendorInvoice"
            WHERE "PurchaseOrderID" = %s
              AND LOWER(TRIM(COALESCE("VendorInvoiceNumber", ''))) = LOWER(TRIM(COALESCE(%s, '')))
            ORDER BY "VendorInvoiceID" DESC
            LIMIT 1
            ''',
            (po_id, invoice_number),
        )
        return cur.fetchone()
    finally:
        conn.close()


def _list_vendor_invoices_for_po(po_id: int, exclude_invoice_id: int | None = None):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        sql = '''
            SELECT
                "VendorInvoiceID",
                "PurchaseOrderID",
                COALESCE("VendorInvoiceNumber", '') AS "VendorInvoiceNumber",
                COALESCE("VendorInvoiceDate"::text, '') AS "VendorInvoiceDate",
                COALESCE("VendorInvoiceDueDate"::text, '') AS "VendorInvoiceDueDate",
                COALESCE("VendorInvoiceAmount", 0) AS "VendorInvoiceAmount",
                COALESCE("VendorInvoiceStatus", '') AS "VendorInvoiceStatus",
                COALESCE("Description", '') AS "Description"
            FROM "VendorInvoice"
            WHERE "PurchaseOrderID" = %s
        '''
        params = [int(po_id)]
        if exclude_invoice_id:
            sql += ' AND "VendorInvoiceID" != %s'
            params.append(int(exclude_invoice_id))
        sql += ' ORDER BY "VendorInvoiceID" DESC'
        cur.execute(sql, params)
        return cur.fetchall()
    finally:
        conn.close()


def _normalize_compare_text(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _coerce_compare_float(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _normalize_line_compare_entries(line_candidates) -> list[dict]:
    normalized = []
    for item in line_candidates or []:
        if not isinstance(item, dict):
            continue
        part_number = _normalize_compare_text(item.get("part_number"))
        description = _normalize_compare_text(item.get("description"))
        quantity = _coerce_compare_float(
            item.get("quantity")
            if item.get("quantity") not in (None, "")
            else item.get("received_qty_candidate")
            if item.get("received_qty_candidate") not in (None, "")
            else item.get("invoice_qty")
            if item.get("invoice_qty") not in (None, "")
            else item.get("invoiced_qty")
        )
        line_amount = _coerce_compare_float(
            item.get("line_amount")
            if item.get("line_amount") not in (None, "")
            else item.get("line_total")
            if item.get("line_total") not in (None, "")
            else item.get("invoice_line_total")
            if item.get("invoice_line_total") not in (None, "")
            else item.get("subtotal")
        )
        if not part_number and not description:
            continue
        normalized.append(
            {
                "part_number": part_number,
                "description": description,
                "quantity": round(float(quantity), 3) if quantity is not None else None,
                "line_amount": round(float(line_amount), 2) if line_amount is not None else None,
            }
        )
    return normalized


def _build_line_compare_counter(entries: list[dict]) -> Counter:
    counter = Counter()
    for item in entries:
        part_number = str(item.get("part_number") or "")
        description = str(item.get("description") or "")
        key = (
            part_number,
            "" if part_number else description,
            item.get("quantity"),
        )
        counter[key] += 1
    return counter


def _summarize_line_similarity(incoming_lines: list[dict], existing_lines: list[dict]) -> dict:
    incoming_counter = _build_line_compare_counter(incoming_lines)
    existing_counter = _build_line_compare_counter(existing_lines)
    overlap_counter = incoming_counter & existing_counter
    overlap_count = int(sum(overlap_counter.values()))
    incoming_count = int(sum(incoming_counter.values()))
    existing_count = int(sum(existing_counter.values()))
    incoming_qty_total = round(
        sum(float(item.get("quantity") or 0.0) for item in incoming_lines if item.get("quantity") is not None),
        3,
    )
    existing_qty_total = round(
        sum(float(item.get("quantity") or 0.0) for item in existing_lines if item.get("quantity") is not None),
        3,
    )
    strong_match = (
        incoming_count > 0
        and existing_count > 0
        and overlap_count == min(incoming_count, existing_count)
        and abs(incoming_qty_total - existing_qty_total) <= 0.001
    )
    return {
        "incoming_line_count": incoming_count,
        "existing_line_count": existing_count,
        "overlapping_line_count": overlap_count,
        "incoming_quantity_total": incoming_qty_total,
        "existing_quantity_total": existing_qty_total,
        "strong_match": strong_match,
    }


def _extract_receipt_identity_from_key(receipt_key: str | None) -> int | None:
    text = str(receipt_key or "").strip()
    if not text.startswith("receipt:"):
        return None
    try:
        return int(text.split(":", 1)[1])
    except (TypeError, ValueError):
        return None


def _resolve_receipt_context_for_duplicate_check(po_id: int, packing_slip_number: str) -> dict:
    normalized_packing_slip = str(packing_slip_number or "").strip()
    result = {
        "receipt_key": None,
        "receipt_id": None,
        "receipt_ref": normalized_packing_slip or None,
        "receipt_choices": [],
    }
    if not normalized_packing_slip:
        return result

    for choice in get_purchase_order_receipt_choices(int(po_id)):
        result["receipt_choices"].append(
            {
                "ReceiptID": choice.get("ReceiptID"),
                "Key": choice.get("Key"),
                "DocumentRef": choice.get("DocumentRef"),
            }
        )
        doc_ref = str(choice.get("DocumentRef") or "").strip()
        if doc_ref and doc_ref.lower() == normalized_packing_slip.lower():
            result["receipt_key"] = str(choice.get("Key") or "").strip() or None
            result["receipt_id"] = int(choice["ReceiptID"]) if choice.get("ReceiptID") not in (None, "") else None
            result["receipt_ref"] = doc_ref
            break
    return result


def analyze_vendor_invoice_duplicate_risk(
    *,
    po_id: int,
    invoice_number: str,
    packing_slip_number: str = "",
    total_amount: float | None = None,
    line_candidates: list | None = None,
    exclude_invoice_id: int | None = None,
):
    normalized_invoice_number = str(invoice_number or "").strip()
    normalized_packing_slip = str(packing_slip_number or "").strip()
    normalized_amount = float(total_amount or 0.0)
    incoming_lines = _normalize_line_compare_entries(line_candidates or [])
    incoming_receipt = _resolve_receipt_context_for_duplicate_check(int(po_id), normalized_packing_slip)

    result = {
        "supported": True,
        "checked": True,
        "duplicate_found": False,
        "duplicate_outcome": None,
        "duplicate_reason": None,
        "review_required": False,
        "review_reason": None,
        "matched_purchase_order_id": int(po_id),
        "matched_vendor_invoice_ids": [],
        "existing_invoice_numbers": [],
        "matching_packing_slip": normalized_packing_slip or None,
        "matching_receipt_ids": [],
        "matching_receipt_refs": [],
        "exact_duplicate_vendor_invoice_id": None,
        "reusable_existing_draft_id": None,
        "amount_comparison": {
            "input_total_amount": round(normalized_amount, 2) if normalized_amount else 0.0,
            "matched_invoice_amounts": [],
        },
        "line_comparison_summary": {
            "input_line_count": len(incoming_lines),
            "strong_match_invoice_ids": [],
            "comparisons": [],
        },
        "uncertainty_notes": [],
        "existing_matches": [],
    }

    if not normalized_invoice_number and not normalized_packing_slip and not normalized_amount and not incoming_lines:
        result["checked"] = False
        result["uncertainty_notes"].append("Not enough invoice identity was supplied for semantic duplicate detection.")
        return result

    exact_duplicate = _find_vendor_invoice_by_number(int(po_id), normalized_invoice_number) if normalized_invoice_number else None
    if exact_duplicate and (not exclude_invoice_id or int(exact_duplicate.get("VendorInvoiceID") or 0) != int(exclude_invoice_id)):
        invoice_id = int(exact_duplicate.get("VendorInvoiceID") or 0)
        existing_status = str(exact_duplicate.get("VendorInvoiceStatus") or "").strip() or "Draft"
        result["duplicate_found"] = True
        result["duplicate_outcome"] = "DUPLICATE_INVOICE_POSSIBLE"
        result["duplicate_reason"] = (
            f"Vendor invoice {normalized_invoice_number} already exists for PO #{po_id}."
        )
        result["exact_duplicate_vendor_invoice_id"] = invoice_id
        result["matched_vendor_invoice_ids"] = [invoice_id]
        result["existing_invoice_numbers"] = [str(exact_duplicate.get("VendorInvoiceNumber") or normalized_invoice_number)]
        if existing_status == "Draft":
            result["reusable_existing_draft_id"] = invoice_id
        existing_amount = _coerce_compare_float(exact_duplicate.get("VendorInvoiceAmount"))
        result["amount_comparison"]["matched_invoice_amounts"].append(
            {
                "vendor_invoice_id": invoice_id,
                "invoice_number": str(exact_duplicate.get("VendorInvoiceNumber") or normalized_invoice_number),
                "invoice_total_amount": round(existing_amount, 2) if existing_amount is not None else None,
                "difference": round(abs(existing_amount - normalized_amount), 2)
                if existing_amount is not None and normalized_amount
                else None,
            }
        )
        return result

    severity_rank = {
        "DUPLICATE_PACKING_SLIP_INVOICE_POSSIBLE": 4,
        "DUPLICATE_RECEIPT_BILLING_POSSIBLE": 3,
        "POSSIBLE_REBILL_OF_ALREADY_INVOICED_PARTS": 2,
        "DUPLICATE_INVOICE_POSSIBLE": 1,
    }
    chosen_match = None
    existing_invoice_count = 0

    for row in _list_vendor_invoices_for_po(int(po_id), exclude_invoice_id=exclude_invoice_id):
        existing_invoice_count += 1
        invoice_id = int(row.get("VendorInvoiceID") or 0)
        memory = get_vendor_invoice_memory(invoice_id)
        existing_number = str(row.get("VendorInvoiceNumber") or "").strip()
        existing_amount = _coerce_compare_float(row.get("VendorInvoiceAmount"))
        existing_packing_slip = str(
            memory.get("packing_slip_number")
            or memory.get("receipt_ref")
            or ""
        ).strip()
        existing_receipt_key = str(memory.get("receipt_key") or "").strip()
        existing_receipt_id = _extract_receipt_identity_from_key(existing_receipt_key)
        existing_receipt_ref = str(memory.get("receipt_ref") or "").strip()
        existing_lines = _normalize_line_compare_entries(memory.get("line_candidates") or memory.get("rows") or [])
        amount_difference = (
            abs(float(existing_amount) - normalized_amount)
            if existing_amount is not None and normalized_amount
            else None
        )
        similar_total = (
            existing_amount is not None
            and normalized_amount > 0
            and amount_difference is not None
            and amount_difference <= max(0.01, min(5.0, max(abs(existing_amount), abs(normalized_amount)) * 0.005))
        )
        line_similarity = _summarize_line_similarity(incoming_lines, existing_lines)
        distinct_packing_slip_evidence = bool(
            normalized_packing_slip
            and existing_packing_slip
            and existing_packing_slip.lower() != normalized_packing_slip.lower()
        )
        distinct_receipt_evidence = bool(
            incoming_receipt.get("receipt_id") is not None
            and existing_receipt_id is not None
            and int(incoming_receipt["receipt_id"]) != int(existing_receipt_id)
        ) or bool(
            incoming_receipt.get("receipt_ref")
            and existing_receipt_ref
            and str(incoming_receipt["receipt_ref"]).strip().lower() != existing_receipt_ref.lower()
        )
        distinct_shipment_evidence = bool(distinct_packing_slip_evidence or distinct_receipt_evidence)

        outcome = None
        reason = None
        if normalized_packing_slip and existing_packing_slip and existing_packing_slip.lower() == normalized_packing_slip.lower():
            outcome = "DUPLICATE_PACKING_SLIP_INVOICE_POSSIBLE"
            reason = (
                f"Packing slip {normalized_packing_slip} is already linked to vendor invoice {existing_number or invoice_id} on PO #{po_id}."
            )
        elif (
            incoming_receipt.get("receipt_id") is not None
            and existing_receipt_id is not None
            and int(incoming_receipt["receipt_id"]) == int(existing_receipt_id)
        ) or (
            incoming_receipt.get("receipt_ref")
            and existing_receipt_ref
            and str(incoming_receipt["receipt_ref"]).strip().lower() == existing_receipt_ref.lower()
        ):
            outcome = "DUPLICATE_RECEIPT_BILLING_POSSIBLE"
            reason = (
                f"Receipt evidence for {incoming_receipt.get('receipt_ref') or normalized_packing_slip or 'this invoice'} "
                f"already appears on vendor invoice {existing_number or invoice_id} for PO #{po_id}."
            )
        elif similar_total and line_similarity.get("strong_match") and not distinct_shipment_evidence:
            outcome = "POSSIBLE_REBILL_OF_ALREADY_INVOICED_PARTS"
            reason = (
                f"Invoice total and matched line quantities look like a rebill of parts already invoiced on PO #{po_id}."
            )

        result["line_comparison_summary"]["comparisons"].append(
            {
                "vendor_invoice_id": invoice_id,
                "invoice_number": existing_number,
                "same_total": bool(similar_total),
                "amount_difference": round(float(amount_difference), 2) if amount_difference is not None else None,
                "overlapping_line_count": int(line_similarity.get("overlapping_line_count") or 0),
                "strong_match": bool(line_similarity.get("strong_match")),
                "distinct_shipment_evidence": bool(distinct_shipment_evidence),
            }
        )
        if line_similarity.get("strong_match"):
            result["line_comparison_summary"]["strong_match_invoice_ids"].append(invoice_id)
        if similar_total:
            result["amount_comparison"]["matched_invoice_amounts"].append(
                {
                    "vendor_invoice_id": invoice_id,
                    "invoice_number": existing_number,
                    "invoice_total_amount": round(float(existing_amount), 2) if existing_amount is not None else None,
                    "difference": round(float(amount_difference), 2) if amount_difference is not None else None,
                }
            )

        if not outcome:
            continue

        result["duplicate_found"] = True
        result["matched_vendor_invoice_ids"].append(invoice_id)
        if existing_number:
            result["existing_invoice_numbers"].append(existing_number)
        if existing_receipt_id is not None:
            result["matching_receipt_ids"].append(existing_receipt_id)
        if existing_receipt_ref:
            result["matching_receipt_refs"].append(existing_receipt_ref)
        result["existing_matches"].append(
            {
                "vendor_invoice_id": invoice_id,
                "invoice_number": existing_number,
                "vendor_invoice_status": str(row.get("VendorInvoiceStatus") or "").strip(),
                "packing_slip_number": existing_packing_slip or None,
                "receipt_id": existing_receipt_id,
                "receipt_ref": existing_receipt_ref or None,
                "same_total": bool(similar_total),
                "line_similarity": line_similarity,
                "distinct_shipment_evidence": bool(distinct_shipment_evidence),
                "duplicate_outcome": outcome,
            }
        )
        current_rank = severity_rank.get(outcome, 0)
        chosen_rank = severity_rank.get(chosen_match["outcome"], 0) if chosen_match else -1
        if current_rank > chosen_rank:
            chosen_match = {
                "outcome": outcome,
                "reason": reason,
            }

    if chosen_match:
        result["duplicate_outcome"] = chosen_match["outcome"]
        result["duplicate_reason"] = chosen_match["reason"]
        if chosen_match["outcome"] == "POSSIBLE_REBILL_OF_ALREADY_INVOICED_PARTS":
            result["uncertainty_notes"].append(
                "Similar total and line quantities matched a prior vendor invoice. Operator review is required before any new draft is staged."
            )

    has_incoming_shipment_identity = bool(
        normalized_packing_slip
        or incoming_receipt.get("receipt_id") is not None
        or str(incoming_receipt.get("receipt_ref") or "").strip()
    )
    has_incoming_line_identity = bool(incoming_lines)
    if (
        not result["duplicate_found"]
        and existing_invoice_count > 0
        and not has_incoming_shipment_identity
        and not has_incoming_line_identity
    ):
        result["review_required"] = True
        result["review_reason"] = (
            f"PO #{po_id} already has vendor invoice history, but this invoice does not include enough shipment evidence "
            "to distinguish it from prior billing. Review before staging a new draft payable record."
        )
        result["uncertainty_notes"].append(
            "Prior invoices exist on this PO, but no packing slip, receipt, or line-level shipment evidence was supplied for this invoice."
        )

    result["matched_vendor_invoice_ids"] = sorted({int(value) for value in result["matched_vendor_invoice_ids"] if value not in (None, "")})
    result["existing_invoice_numbers"] = sorted({str(value) for value in result["existing_invoice_numbers"] if str(value or "").strip()})
    result["matching_receipt_ids"] = sorted({int(value) for value in result["matching_receipt_ids"] if value not in (None, "")})
    result["matching_receipt_refs"] = sorted({str(value) for value in result["matching_receipt_refs"] if str(value or "").strip()})
    return result


def _next_vendor_invoice_number(po_id: int):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT COUNT(*) AS "InvoiceCount"
            FROM "VendorInvoice"
            WHERE "PurchaseOrderID" = %s
            ''',
            (po_id,),
        )
        row = cur.fetchone() or {}
    finally:
        conn.close()
    return f"DRAFT-PO-{po_id}-{int(row.get('InvoiceCount') or 0) + 1}"


def _upsert_vendor_invoice_row(invoice_id: int = None, po_id: int = None, invoice_number: str = "", invoice_date: str = "", due_date: str = "", amount: float = 0.0, description: str = "", status: str = "Draft"):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        if invoice_id:
            cur.execute(
                '''
                UPDATE "VendorInvoice"
                SET "VendorInvoiceNumber" = %s,
                    "VendorInvoiceDate" = %s,
                    "VendorInvoiceDueDate" = %s,
                    "VendorInvoiceAmount" = %s,
                    "VendorInvoiceStatus" = %s,
                    "Description" = %s
                WHERE "VendorInvoiceID" = %s
                RETURNING *
                ''',
                (invoice_number, invoice_date, due_date or None, amount, status, description, invoice_id),
            )
            row = cur.fetchone()
        else:
            cur.execute(
                '''
                INSERT INTO "VendorInvoice"
                ("VendorInvoiceNumber", "PurchaseOrderID", "VendorInvoiceDate", "VendorInvoiceDueDate", "VendorInvoiceAmount", "VendorInvoiceStatus", "Description")
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                RETURNING *
                ''',
                (invoice_number or f"DRAFT-PO-{po_id}", po_id, invoice_date or datetime.date.today().isoformat(), due_date or None, amount, status, description),
            )
            row = cur.fetchone()
        conn.commit()
        return row
    finally:
        conn.close()


def get_purchase_order_invoice_rollup(po_id: int, exclude_invoice_id: int = None):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT COALESCE("PurchaseOrderTotal", 0) AS "PurchaseOrderTotal"
            FROM "PurchaseOrder"
            WHERE "PurchaseOrderID" = %s
            ''',
            (po_id,),
        )
        po_row = cur.fetchone() or {}

        sql = '''
            SELECT COALESCE(SUM(COALESCE("VendorInvoiceAmount", 0)), 0) AS "InvoicedTotal"
            FROM "VendorInvoice"
            WHERE "PurchaseOrderID" = %s
        '''
        params = [po_id]
        if exclude_invoice_id:
            sql += ' AND "VendorInvoiceID" != %s'
            params.append(exclude_invoice_id)
        cur.execute(sql, params)
        invoice_row = cur.fetchone() or {}
    finally:
        conn.close()

    po_total = float(po_row.get("PurchaseOrderTotal") or 0)
    invoiced_total = float(invoice_row.get("InvoicedTotal") or 0)
    return {
        "po_total": po_total,
        "invoiced_total": invoiced_total,
        "remaining_total": po_total - invoiced_total,
    }


def _get_vendor_invoice_item_rollup(po_id: int, exclude_invoice_id: int = None):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        sql = '''
            SELECT "VendorInvoiceID"
            FROM "VendorInvoice"
            WHERE "PurchaseOrderID" = %s
        '''
        params = [po_id]
        if exclude_invoice_id:
            sql += ' AND "VendorInvoiceID" != %s'
            params.append(exclude_invoice_id)
        cur.execute(sql, params)
        invoice_rows = cur.fetchall()
    finally:
        conn.close()

    totals = {}
    for invoice_row in invoice_rows:
        memory = get_vendor_invoice_memory(int(invoice_row["VendorInvoiceID"]))
        for row in memory.get("rows", []):
            po_item_id = int(row.get("po_item_id") or 0)
            if not po_item_id:
                continue
            totals[po_item_id] = totals.get(po_item_id, 0.0) + float(row.get("invoiced_qty") or 0)
    return totals


def _resolve_receipt_choice(po_id: int, receipt_key: str = None):
    choices = get_purchase_order_receipt_choices(po_id)
    choice_map = {row["Key"]: row for row in choices}
    if receipt_key and receipt_key in choice_map:
        return choice_map[receipt_key]
    for row in choices:
        if row["Key"] != "all_received":
            return row
    return choice_map.get("all_received")


def create_vendor_invoice_draft(po_id: int, receipt_key: str = None):
    choice = _resolve_receipt_choice(po_id, receipt_key)
    invoice = _upsert_vendor_invoice_row(
        po_id=po_id,
        invoice_number=_next_vendor_invoice_number(po_id),
        invoice_date=datetime.date.today().isoformat(),
        amount=0.0,
        description=f"Draft vendor invoice created from {choice.get('Label') if choice else 'Payables workspace'}.",
        status="Draft",
    )
    rows = _build_default_invoice_rows(po_id, invoice["VendorInvoiceID"], choice["Key"] if choice else "all_received")
    upsert_vendor_invoice_memory(
        invoice["VendorInvoiceID"],
        {
            "po_id": po_id,
            "receipt_key": choice["Key"] if choice else "all_received",
            "receipt_label": choice.get("Label") if choice else "All received material on this PO",
            "receipt_ref": choice.get("DocumentRef") if choice else "All received material",
            "receipt_date": choice.get("ReceiveDate") if choice else "",
            "rows": rows,
            "last_updated_at": datetime.datetime.now().isoformat(),
        },
    )
    return invoice


def get_or_create_vendor_invoice_draft(po_id: int):
    existing = _get_latest_vendor_invoice_for_po(po_id, statuses=["Draft"])
    if existing:
        return existing
    return create_vendor_invoice_draft(po_id, receipt_key="all_received")


def _build_default_invoice_rows(po_id: int, invoice_id: int = None, receipt_key: str = None):
    memory = get_vendor_invoice_memory(invoice_id) if invoice_id else {}
    active_receipt_key = receipt_key or memory.get("receipt_key") or "all_received"
    receipt_qty_map = get_purchase_order_receipt_quantities(po_id, active_receipt_key)
    prior_invoiced_qty_map = _get_vendor_invoice_item_rollup(po_id, exclude_invoice_id=invoice_id)
    line_map = {row.get("po_item_id"): row for row in memory.get("rows", memory.get("parsed_lines", []))}
    rows = []
    for po_item in _get_po_item_rows(po_id):
        saved = line_map.get(po_item["POItemID"], {})
        delivered_qty = float(receipt_qty_map.get(int(po_item["POItemID"]), 0.0))
        if active_receipt_key == "all_received" and not receipt_qty_map:
            delivered_qty = float(po_item.get("QuantityReceived") or 0)
        previously_invoiced_qty = float(prior_invoiced_qty_map.get(int(po_item["POItemID"]), 0.0))
        default_invoice_qty = max(0.0, delivered_qty - previously_invoiced_qty)
        invoiced_qty = float(saved.get("invoice_qty", default_invoice_qty))
        invoice_price = float(saved.get("invoice_unit_price", po_item.get("POUnitPrice") or 0))
        subtotal = float(saved.get("invoice_line_total", invoiced_qty * invoice_price))
        match_status = "OK"
        if invoiced_qty > delivered_qty + 0.001:
            match_status = "Review"
        if abs(invoice_price - float(po_item.get("POUnitPrice") or 0)) > 0.009:
            match_status = "Review"
        rows.append(
            {
                "po_item_id": po_item["POItemID"],
                "description": po_item.get("Description") or "",
                "estimate_unit_price": float(po_item.get("EstimateUnitPrice") or 0),
                "delivered_qty": delivered_qty,
                "invoiced_qty": invoiced_qty,
                "invoice_unit_price": invoice_price,
                "subtotal": subtotal,
                "backordered_qty": float(po_item.get("BackorderedQty") or 0),
                "match_status": match_status,
            }
        )
    return rows


def get_open_purchase_order_invoice_pipeline():
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT
                po."PurchaseOrderID",
                po."Date",
                po."Status",
                po."PurchaseOrderTotal",
                v."VendorName",
                s."SiteName",
                (
                    SELECT vi."VendorInvoiceID"
                    FROM "VendorInvoice" vi
                    WHERE vi."PurchaseOrderID" = po."PurchaseOrderID"
                    ORDER BY vi."VendorInvoiceID" DESC
                    LIMIT 1
                ) AS "VendorInvoiceID",
                (
                    SELECT vi."VendorInvoiceStatus"
                    FROM "VendorInvoice" vi
                    WHERE vi."PurchaseOrderID" = po."PurchaseOrderID"
                    ORDER BY vi."VendorInvoiceID" DESC
                    LIMIT 1
                ) AS "VendorInvoiceStatus",
                (
                    SELECT vi."VendorInvoiceAmount"
                    FROM "VendorInvoice" vi
                    WHERE vi."PurchaseOrderID" = po."PurchaseOrderID"
                    ORDER BY vi."VendorInvoiceID" DESC
                    LIMIT 1
                ) AS "VendorInvoiceAmount"
            FROM "PurchaseOrder" po
            JOIN "Vendor" v ON po."VendorID" = v."VendorID"
            LEFT JOIN "WorkOrder" wo ON po."WorkOrderID" = wo."WorkOrderID"
            LEFT JOIN "Site" s ON wo."SiteID" = s."SiteID"
            WHERE po."Date" >= %s
              AND po."Status" != 'Retired'
              AND (
                    EXISTS (
                        SELECT 1
                        FROM "PurchaseOrderItem" poi
                        WHERE poi."PurchaseOrderID" = po."PurchaseOrderID"
                          AND COALESCE(poi."QuantityReceived", 0) > 0
                    )
                    OR EXISTS (
                        SELECT 1
                        FROM "VendorInvoice" vi
                        WHERE vi."PurchaseOrderID" = po."PurchaseOrderID"
                    )
              )
            ORDER BY po."Date" DESC, po."PurchaseOrderID" DESC
            ''',
            (AUTOMATION_MIN_DATE,),
        )
        return cur.fetchall()
    finally:
        conn.close()


def get_vendor_invoice_pipeline():
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT
                vi."VendorInvoiceID",
                vi."PurchaseOrderID",
                vi."VendorInvoiceNumber",
                vi."VendorInvoiceDate",
                vi."VendorInvoiceDueDate",
                vi."VendorInvoiceAmount",
                vi."VendorInvoiceStatus",
                po."Status" AS "POStatus",
                v."VendorName",
                s."SiteName"
            FROM "VendorInvoice" vi
            JOIN "PurchaseOrder" po ON vi."PurchaseOrderID" = po."PurchaseOrderID"
            JOIN "Vendor" v ON po."VendorID" = v."VendorID"
            LEFT JOIN "WorkOrder" wo ON po."WorkOrderID" = wo."WorkOrderID"
            LEFT JOIN "Site" s ON wo."SiteID" = s."SiteID"
            WHERE po."Status" != 'Retired'
            ORDER BY COALESCE(vi."VendorInvoiceDate"::text, '') DESC, vi."VendorInvoiceID" DESC
            '''
        )
        return cur.fetchall()
    finally:
        conn.close()


def get_received_purchase_order_choices():
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT DISTINCT
                po."PurchaseOrderID",
                v."VendorName",
                COALESCE(s."SiteName", 'Unknown Site') AS "SiteName"
            FROM "PurchaseOrder" po
            JOIN "Vendor" v ON po."VendorID" = v."VendorID"
            LEFT JOIN "WorkOrder" wo ON po."WorkOrderID" = wo."WorkOrderID"
            LEFT JOIN "Site" s ON wo."SiteID" = s."SiteID"
            LEFT JOIN "PurchaseOrderItem" poi ON poi."PurchaseOrderID" = po."PurchaseOrderID"
            WHERE po."Status" != 'Retired'
              AND (
                    COALESCE(poi."QuantityReceived", 0) > 0
                    OR EXISTS (
                        SELECT 1
                        FROM "VendorInvoice" vi
                        WHERE vi."PurchaseOrderID" = po."PurchaseOrderID"
                    )
              )
            ORDER BY po."PurchaseOrderID" DESC
            '''
        )
        rows = cur.fetchall()
        return [
            {
                "PurchaseOrderID": int(row["PurchaseOrderID"]),
                "Label": f"PO #{row['PurchaseOrderID']} - {row.get('VendorName') or 'Vendor'} - {row.get('SiteName') or 'Site'}",
            }
            for row in rows
        ]
    finally:
        conn.close()


def get_vendor_invoice_creation_context(po_id: int):
    po_data = get_po_export_data(po_id)
    if not po_data:
        return None
    return {
        "po": po_data,
        "receipt_choices": get_purchase_order_receipt_choices(po_id),
        "rollup": get_purchase_order_invoice_rollup(po_id),
    }


def get_vendor_invoice_workspace(po_id: int):
    po_data = get_po_export_data(po_id)
    if not po_data:
        return None
    invoice = get_or_create_vendor_invoice_draft(po_id)
    memory = get_vendor_invoice_memory(invoice["VendorInvoiceID"])
    rows = _build_default_invoice_rows(po_id, invoice["VendorInvoiceID"], memory.get("receipt_key"))
    total = sum(float(row.get("subtotal") or 0) for row in rows)
    rollup_excluding_current = get_purchase_order_invoice_rollup(po_id, exclude_invoice_id=invoice["VendorInvoiceID"])
    return {
        "po": po_data,
        "invoice": invoice,
        "rows": rows,
        "total": total,
        "report_path": memory.get("report_path"),
        "receipt_key": memory.get("receipt_key") or "all_received",
        "receipt_label": memory.get("receipt_label") or "All received material on this PO",
        "rollup_excluding_current": rollup_excluding_current,
        "receipt_choices": get_purchase_order_receipt_choices(po_id),
    }


def get_vendor_invoice_workspace_by_invoice_id(invoice_id: int):
    started_at = time.perf_counter()
    try:
        payload = get_vendor_invoice_detail(invoice_id)
        if not payload:
            return None
        po_id = int(payload["header"]["PurchaseOrderID"])
        po_data = get_po_export_data(po_id)
        memory = payload["memory"]
        rows = _build_default_invoice_rows(po_id, invoice_id, memory.get("receipt_key"))
        total = sum(float(row.get("subtotal") or 0) for row in rows)
        return {
            "po": po_data,
            "invoice": payload["header"],
            "rows": rows,
            "total": total,
            "report_path": memory.get("report_path"),
            "receipt_key": memory.get("receipt_key") or "all_received",
            "receipt_label": memory.get("receipt_label") or "All received material on this PO",
            "rollup_excluding_current": get_purchase_order_invoice_rollup(po_id, exclude_invoice_id=invoice_id),
            "receipt_choices": get_purchase_order_receipt_choices(po_id),
        }
    finally:
        _perf_log("db", "vendor_invoices.get_vendor_invoice_workspace_by_invoice_id", started_at)


def _recalculate_purchase_order_total(po_id: int):
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            UPDATE "PurchaseOrder"
            SET "PurchaseOrderTotal" = COALESCE((
                SELECT SUM(COALESCE("LineTotal", 0))
                FROM "PurchaseOrderItem"
                WHERE "PurchaseOrderID" = %s
            ), 0)
            WHERE "PurchaseOrderID" = %s
            ''',
            (po_id, po_id),
        )
        conn.commit()
    finally:
        conn.close()


def _get_po_item(po_item_id: int):
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT poi.*, po."PurchaseOrderID", po."WorkOrderID", wo."SourceEstimateID" AS "EstimateID"
            FROM "PurchaseOrderItem" poi
            JOIN "PurchaseOrder" po ON poi."PurchaseOrderID" = po."PurchaseOrderID"
            LEFT JOIN "WorkOrder" wo ON po."WorkOrderID" = wo."WorkOrderID"
            WHERE poi."POItemID" = %s
            ''',
            (po_item_id,),
        )
        return cur.fetchone()
    finally:
        conn.close()


def apply_vendor_invoice_adjustment(invoice_id: int, po_item_id: int, field: str, new_value: float, note_text: str):
    from neon_ai.database.automation import log_estimate_action

    po_item = _get_po_item(po_item_id)
    if not po_item:
        raise ValueError("Purchase order item not found.")
    if not note_text or not note_text.strip():
        raise ValueError("Adjustment note is required.")

    new_value = float(new_value)
    conn = get_connection()
    cur = conn.cursor()
    try:
        if field == "qty":
            new_qty = max(0.0, new_value)
            cur.execute(
                '''
                UPDATE "PurchaseOrderItem"
                SET "QuantityReceived" = %s,
                    "LastReceivedDate" = %s,
                    "LineTotal" = %s * COALESCE("UnitPriceAtOrder", 0),
                    "LastPackingSlip" = CASE
                        WHEN COALESCE("LastPackingSlip", '') = '' THEN %s
                        ELSE CONCAT("LastPackingSlip", ' | ', %s)
                    END
                WHERE "POItemID" = %s
                ''',
                (
                    new_qty,
                    datetime.date.today().isoformat(),
                    new_qty,
                    "Adjusted via vendor invoice",
                    "Adjusted via vendor invoice",
                    po_item_id,
                ),
            )
        elif field == "price":
            new_price = max(0.0, new_value)
            cur.execute(
                '''
                UPDATE "PurchaseOrderItem"
                SET "UnitPriceAtOrder" = %s,
                    "LineTotal" = COALESCE("QuantityReceived", 0) * %s
                WHERE "POItemID" = %s
                ''',
                (new_price, new_price, po_item_id),
            )
        else:
            raise ValueError("Unsupported adjustment field.")
        conn.commit()
    finally:
        conn.close()

    _recalculate_purchase_order_total(po_item["PurchaseOrderID"])
    update_purchase_order_receiving_status(po_item["PurchaseOrderID"])

    memory = get_vendor_invoice_memory(invoice_id)
    rows = memory.get("rows", [])
    for row in rows:
        if row.get("po_item_id") != po_item_id:
            continue
        if field == "qty":
            row["delivered_qty"] = max(0.0, new_value)
            row["invoiced_qty"] = max(0.0, new_value)
        else:
            row["invoice_unit_price"] = max(0.0, new_value)
        row["subtotal"] = float(row.get("invoiced_qty") or 0) * float(row.get("invoice_unit_price") or 0)
    if rows:
        upsert_vendor_invoice_memory(invoice_id, {"rows": rows, "last_updated_at": datetime.datetime.now().isoformat()})

    estimate_id = po_item.get("EstimateID")
    if estimate_id:
        label = "quantity" if field == "qty" else "price"
        log_estimate_action(
            estimate_id,
            f"Vendor invoice adjustment on PO #{po_item['PurchaseOrderID']} for item '{po_item.get('Description')}'. Updated {label} to {new_value}. Note: {note_text}",
            "Vendor Invoice",
        )


def save_vendor_invoice_draft(
    po_id: int,
    invoice_id: int,
    invoice_number: str,
    invoice_date: str,
    due_date: str,
    description: str,
    rows: list,
    status: str = "Draft",
    receipt_key: str = None,
    receipt_label: str = None,
):
    total = sum(float(row.get("subtotal") or 0) for row in rows)
    invoice = _upsert_vendor_invoice_row(
        invoice_id=invoice_id,
        invoice_number=invoice_number or _next_vendor_invoice_number(po_id),
        invoice_date=invoice_date or datetime.date.today().isoformat(),
        due_date=due_date,
        amount=total,
        description=description or "Vendor invoice draft saved from Payables workspace.",
        status=status,
    )
    existing_memory = get_vendor_invoice_memory(invoice["VendorInvoiceID"])
    choice = _resolve_receipt_choice(po_id, receipt_key or existing_memory.get("receipt_key"))
    resolved_receipt_ref = choice.get("DocumentRef") if choice else ""
    resolved_receipt_date = choice.get("ReceiveDate") if choice else ""
    if not receipt_key:
        resolved_receipt_ref = existing_memory.get("receipt_ref") or resolved_receipt_ref
        resolved_receipt_date = existing_memory.get("receipt_date") or resolved_receipt_date
    upsert_vendor_invoice_memory(
        invoice["VendorInvoiceID"],
        {
            "po_id": po_id,
            "rows": rows,
            "receipt_key": receipt_key or existing_memory.get("receipt_key") or "all_received",
            "receipt_label": receipt_label or existing_memory.get("receipt_label") or (choice.get("Label") if choice else "All received material on this PO"),
            "receipt_ref": resolved_receipt_ref,
            "receipt_date": resolved_receipt_date,
            "last_updated_at": datetime.datetime.now().isoformat(),
        },
    )
    return invoice


def create_vendor_invoice_intake_draft(
    *,
    po_id: int,
    invoice_number: str,
    invoice_date: str = "",
    due_date: str = "",
    total_amount: float | None = None,
    description: str = "",
    vendor_name: str = "",
    packing_slip_number: str = "",
    line_candidates: list | None = None,
    source_proposal_id: int | None = None,
    source_inbound_message_id: int | None = None,
    applied_by: str = "automation_apply",
):
    ensure_vendor_invoice_schema()
    po_header = get_po_export_data(int(po_id))
    if not po_header:
        raise ValueError(f"PurchaseOrder #{po_id} could not be found.")

    normalized_invoice_number = str(invoice_number or "").strip()
    if not normalized_invoice_number:
        raise ValueError("Vendor invoice draft intake requires an invoice number.")

    normalized_vendor_name = str(vendor_name or "").strip()
    normalized_packing_slip = str(packing_slip_number or "").strip()
    normalized_description = str(description or "").strip() or "Vendor invoice draft created from approved automation intake."
    normalized_amount = float(total_amount or 0.0)

    duplicate_result = analyze_vendor_invoice_duplicate_risk(
        po_id=int(po_id),
        invoice_number=normalized_invoice_number,
        packing_slip_number=normalized_packing_slip,
        total_amount=normalized_amount,
        line_candidates=line_candidates or [],
    )
    if (duplicate_result.get("duplicate_found") or duplicate_result.get("review_required")) and not duplicate_result.get("reusable_existing_draft_id"):
        raise ValueError(
            str(
                duplicate_result.get("duplicate_reason")
                or duplicate_result.get("review_reason")
                or f"Possible duplicate vendor invoice intake detected for PO #{po_id}."
            )
        )

    existing = _find_vendor_invoice_by_number(int(po_id), normalized_invoice_number)
    draft_reused = False
    if existing:
        current_status = str(existing.get("VendorInvoiceStatus") or "").strip() or "Draft"
        if current_status != "Draft":
            raise ValueError(
                f"Vendor invoice {normalized_invoice_number} already exists for PO #{po_id} with status {current_status}."
            )
        existing_amount = float(existing.get("VendorInvoiceAmount") or 0.0)
        existing_invoice_date = str(existing.get("VendorInvoiceDate") or "").strip()
        existing_due_date = str(existing.get("VendorInvoiceDueDate") or "").strip()
        if normalized_amount and existing_amount and abs(existing_amount - normalized_amount) > 0.01:
            raise ValueError(
                f"Vendor invoice {normalized_invoice_number} already exists as Draft for PO #{po_id} with a different amount."
            )
        if invoice_date and existing_invoice_date and existing_invoice_date != invoice_date:
            raise ValueError(
                f"Vendor invoice {normalized_invoice_number} already exists as Draft for PO #{po_id} with a different invoice date."
            )
        if due_date and existing_due_date and existing_due_date != due_date:
            raise ValueError(
                f"Vendor invoice {normalized_invoice_number} already exists as Draft for PO #{po_id} with a different due date."
            )
        draft_reused = True
        invoice_row = _upsert_vendor_invoice_row(
            invoice_id=int(existing["VendorInvoiceID"]),
            invoice_number=normalized_invoice_number,
            invoice_date=invoice_date or existing_invoice_date or datetime.date.today().isoformat(),
            due_date=due_date or existing_due_date,
            amount=normalized_amount or existing_amount,
            description=normalized_description or str(existing.get("Description") or "").strip(),
            status="Draft",
        )
    else:
        invoice_row = _upsert_vendor_invoice_row(
            po_id=int(po_id),
            invoice_number=normalized_invoice_number,
            invoice_date=invoice_date or datetime.date.today().isoformat(),
            due_date=due_date,
            amount=normalized_amount,
            description=normalized_description,
            status="Draft",
        )

    invoice_id = int(invoice_row["VendorInvoiceID"])
    existing_memory = get_vendor_invoice_memory(invoice_id)
    receipt_key = existing_memory.get("receipt_key") or "all_received"
    receipt_label = existing_memory.get("receipt_label") or "All received material on this PO"
    receipt_ref = existing_memory.get("receipt_ref") or ""
    receipt_date = existing_memory.get("receipt_date") or ""

    if normalized_packing_slip:
        for choice in get_purchase_order_receipt_choices(int(po_id)):
            doc_ref = str(choice.get("DocumentRef") or "").strip()
            if doc_ref and doc_ref.lower() == normalized_packing_slip.lower():
                receipt_key = str(choice.get("Key") or receipt_key)
                receipt_label = str(choice.get("Label") or receipt_label)
                receipt_ref = doc_ref
                receipt_date = str(choice.get("ReceiveDate") or receipt_date)
                break

    default_rows = _build_default_invoice_rows(int(po_id), invoice_id, receipt_key)
    default_row_total = round(sum(float(row.get("subtotal") or 0.0) for row in default_rows), 2)
    warnings = []
    if normalized_amount and abs(default_row_total - normalized_amount) > 0.01:
        warnings.append(
            "Extracted invoice total differs from PO/receipt-based default draft rows. Operator reconciliation is still required."
        )

    memory_payload = dict(existing_memory or {})
    memory_payload.update(
        {
            "po_id": int(po_id),
            "rows": default_rows,
            "receipt_key": receipt_key,
            "receipt_label": receipt_label,
            "receipt_ref": receipt_ref,
            "receipt_date": receipt_date,
            "source_proposal_id": source_proposal_id,
            "source_inbound_message_id": source_inbound_message_id,
            "draft_origin": "automation_apply_vendor_invoice_intake",
            "applied_by": applied_by,
            "invoice_number": normalized_invoice_number,
            "invoice_date": invoice_date or str(invoice_row.get("VendorInvoiceDate") or "").strip(),
            "due_date": due_date or str(invoice_row.get("VendorInvoiceDueDate") or "").strip(),
            "invoice_total_amount": normalized_amount,
            "vendor_name": normalized_vendor_name or str(po_header.get("VendorName") or "").strip(),
            "packing_slip_number": normalized_packing_slip or None,
            "line_candidates": line_candidates or [],
            "default_row_total": default_row_total,
            "warnings": warnings,
            "duplicate_check_result": duplicate_result,
            "last_updated_at": datetime.datetime.now().isoformat(),
        }
    )
    upsert_vendor_invoice_memory(invoice_id, memory_payload)
    return {
        "invoice": invoice_row,
        "draft_reused": draft_reused,
        "po_id": int(po_id),
        "default_row_total": default_row_total,
        "warnings": warnings,
        "duplicate_result": duplicate_result,
        "memory": memory_payload,
    }


def generate_vendor_invoice_report(invoice_id: int, force: bool = False):
    payload = get_vendor_invoice_detail(invoice_id)
    if not payload:
        raise ValueError("Vendor invoice could not be loaded.")

    memory = get_vendor_invoice_memory(invoice_id)
    existing_path = memory.get("report_path")
    if existing_path and os.path.exists(existing_path) and not force:
        return existing_path

    header = payload["header"]
    po_data = get_po_export_data(header["PurchaseOrderID"])
    target_dir = get_target_folder(
        po_data.get("StreetNumber"),
        po_data.get("StreetName"),
        po_data.get("SiteName"),
        po_data.get("EstimateID"),
        "Vendor Invoices",
    )
    filename = f"Vendor_Invoice_Reconciliation_{invoice_id}.docx"
    report_path = os.path.join(target_dir, filename).replace("\\", "/")

    doc = Document()
    doc.add_heading(f"Vendor Invoice Reconciliation #{invoice_id}", 0)
    doc.add_paragraph(f"Vendor: {header.get('VendorName')}")
    doc.add_paragraph(f"PO #: {header.get('PurchaseOrderID')}")
    doc.add_paragraph(f"Invoice #: {header.get('VendorInvoiceNumber')}")
    doc.add_paragraph(f"Invoice Date: {header.get('VendorInvoiceDate')}")
    doc.add_paragraph(f"Status: {header.get('VendorInvoiceStatus')}")
    if memory.get("receipt_label"):
        doc.add_paragraph(f"Receiving paperwork: {memory.get('receipt_label')}")

    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Description", "Estimate Price", "Delivered Qty", "Invoiced Qty", "Invoice Price", "Subtotal", "Match"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    total = 0.0
    for row in payload["rows"]:
        cells = table.add_row().cells
        cells[0].text = str(row.get("Description") or "")
        cells[1].text = f"${float(row.get('EstimateUnitPrice') or 0):,.2f}"
        cells[2].text = f"{float(row.get('DeliveredQty') or 0):.2f}"
        cells[3].text = f"{float(row.get('InvoicedQty') or 0):.2f}"
        cells[4].text = f"${float(row.get('InvoiceUnitPrice') or 0):,.2f}"
        cells[5].text = f"${float(row.get('InvoiceLineTotal') or 0):,.2f}"
        cells[6].text = str(row.get("MatchStatus") or "OK")
        total += float(row.get("InvoiceLineTotal") or 0)

    doc.add_paragraph(f"Total Invoiced: ${total:,.2f}")
    doc.save(report_path)
    upsert_vendor_invoice_memory(invoice_id, {"report_path": report_path, "last_report_generated_at": datetime.datetime.now().isoformat()})
    return report_path


def get_vendor_invoice_detail(invoice_id: int):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT
                vi.*,
                po."PurchaseOrderID",
                po."Status" AS "POStatus",
                v."VendorName",
                s."SiteName"
            FROM "VendorInvoice" vi
            JOIN "PurchaseOrder" po ON vi."PurchaseOrderID" = po."PurchaseOrderID"
            JOIN "Vendor" v ON po."VendorID" = v."VendorID"
            LEFT JOIN "WorkOrder" wo ON po."WorkOrderID" = wo."WorkOrderID"
            LEFT JOIN "Site" s ON wo."SiteID" = s."SiteID"
            WHERE vi."VendorInvoiceID" = %s
            ''',
            (invoice_id,),
        )
        header = cur.fetchone()
    finally:
        conn.close()

    if not header:
        return None

    memory = get_vendor_invoice_memory(invoice_id)
    rows = _build_default_invoice_rows(header["PurchaseOrderID"], invoice_id, memory.get("receipt_key"))
    return {
        "header": header,
        "rows": [
            {
                "POItemID": row["po_item_id"],
                "Description": row["description"],
                "EstimateUnitPrice": row["estimate_unit_price"],
                "DeliveredQty": row["delivered_qty"],
                "InvoicedQty": row["invoiced_qty"],
                "InvoiceUnitPrice": row["invoice_unit_price"],
                "InvoiceLineTotal": row["subtotal"],
                "MatchStatus": row["match_status"],
                "BackorderedQty": row["backordered_qty"],
            }
            for row in rows
        ],
        "memory": memory,
    }


def process_inbound_vendor_invoice_email(po_id: int, sender_email: str, subject: str, body: str, attachment_path: str = None, received_at=None):
    """LEGACY_AUTOMATION_DISABLED_BY_DEFAULT: processes old inbound vendor invoice email automation only when explicitly enabled."""
    if not require_legacy_automation_runtime("database.vendor_invoices.process_inbound_vendor_invoice_email"):
        return {"skipped": True, "reason": "legacy_automation_disabled"}

    from neon_ai.gateway import send_to_user

    reference_date = received_at.date() if hasattr(received_at, "date") else datetime.date.today()
    attachment_text = ""
    if attachment_path and os.path.exists(attachment_path):
        attachment_text = extract_text_from_attachment(attachment_path)
        if not attachment_text and attachment_path.lower().endswith(".pdf"):
            attachment_text = ocr_pdf_attachment(attachment_path)

    combined_text = "\n".join(part for part in [subject, body, attachment_text] if part)
    header = extract_invoice_header_fields(subject, body, attachment_text, reference_date=reference_date)
    parsed_lines = parse_vendor_invoice_lines(po_id, combined_text)

    existing_invoice = _find_vendor_invoice_by_number(po_id, header["invoice_number"])
    invoice = existing_invoice or create_vendor_invoice_draft(po_id, receipt_key="all_received")
    rows = _build_default_invoice_rows(po_id, invoice["VendorInvoiceID"], get_vendor_invoice_memory(invoice["VendorInvoiceID"]).get("receipt_key"))
    if parsed_lines:
        line_map = {row["po_item_id"]: row for row in parsed_lines}
        for row in rows:
            matched = line_map.get(row["po_item_id"])
            if matched:
                row["invoiced_qty"] = float(matched.get("invoice_qty") or row["delivered_qty"])
                row["invoice_unit_price"] = float(matched.get("invoice_unit_price") or row["invoice_unit_price"])
                row["subtotal"] = row["invoiced_qty"] * row["invoice_unit_price"]
                row["match_status"] = "Review" if row["invoiced_qty"] > row["delivered_qty"] + 0.001 else "OK"

    saved_invoice = save_vendor_invoice_draft(
        po_id,
        invoice["VendorInvoiceID"],
        header["invoice_number"] or invoice.get("VendorInvoiceNumber"),
        header["invoice_date"],
        "",
        f"Draft created from vendor email from {sender_email}.",
        rows,
        status="Draft",
    )
    upsert_vendor_invoice_memory(
        saved_invoice["VendorInvoiceID"],
        {
            "attachment_path": attachment_path,
            "attachment_text_excerpt": (attachment_text or "")[:2000],
            "parsed_lines": parsed_lines,
            "received_at": received_at.isoformat() if hasattr(received_at, "isoformat") else datetime.datetime.now().isoformat(),
            "last_updated_at": datetime.datetime.now().isoformat(),
        },
    )

    po_data = get_po_export_data(po_id)
    send_to_user(
        subject=f"Vendor Invoice Ready For Review: PO #{po_id}",
        content=(
            f"Vendor invoice draft #{saved_invoice['VendorInvoiceID']} is ready for review.\n\n"
            f"Vendor: {po_data.get('VendorName') if po_data else 'Unknown'}\n"
            f"PO: #{po_id}\n"
            f"Invoice number: {header['invoice_number'] or saved_invoice['VendorInvoiceNumber']}\n"
            f"Invoice date: {header['invoice_date']}\n"
            f"Parsed line matches: {len(parsed_lines)}\n\n"
            "Next action: review the draft vendor invoice in the Payables workspace and lock it once verified."
        )
    )

    return {
        "vendor_invoice_id": saved_invoice["VendorInvoiceID"],
        "po_id": po_id,
        "invoice_number": saved_invoice["VendorInvoiceNumber"],
        "invoice_amount": float(saved_invoice.get("VendorInvoiceAmount") or 0),
        "matched_lines": len(parsed_lines),
    }


def lock_vendor_invoice_ready_to_pay(invoice_id: int):
    from neon_ai.gateway import MY_EMAIL, send_to_user

    payload = get_vendor_invoice_detail(invoice_id)
    if not payload:
        raise ValueError("Vendor invoice could not be loaded.")

    header = payload["header"]
    report_path = generate_vendor_invoice_report(invoice_id, force=True)
    _upsert_vendor_invoice_row(
        invoice_id=invoice_id,
        invoice_number=header["VendorInvoiceNumber"],
        invoice_date=header.get("VendorInvoiceDate") or datetime.date.today().isoformat(),
        due_date=header.get("VendorInvoiceDueDate") or "",
        amount=float(header.get("VendorInvoiceAmount") or 0),
        description=header.get("Description") or "",
        status="ReadyToPay",
    )
    upsert_vendor_invoice_memory(invoice_id, {"report_path": report_path, "locked_at": datetime.datetime.now().isoformat()})

    send_to_user(
        subject=f"Pay Vendor Invoice: {header['VendorInvoiceNumber']}",
        content=(
            f"Vendor invoice {header['VendorInvoiceNumber']} for PO #{header['PurchaseOrderID']} is locked and ready to pay.\n\n"
            f"Vendor: {header.get('VendorName')}\n"
            f"Amount: ${float(header.get('VendorInvoiceAmount') or 0):,.2f}\n"
            f"Site: {header.get('SiteName') or 'N/A'}\n\n"
            "The reconciliation report is attached."
        ),
        recipient=MY_EMAIL,
        attachment_path=report_path,
    )
    return report_path


def mark_vendor_invoice_paid(invoice_id: int):
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            '''
            UPDATE "VendorInvoice"
            SET "VendorInvoiceStatus" = 'Paid',
                "VendorInvoicePaidDate" = CURRENT_DATE
            WHERE "VendorInvoiceID" = %s
            ''',
            (invoice_id,),
        )
        conn.commit()
    finally:
        conn.close()


def get_vendor_payables_summary():
    ensure_vendor_invoice_schema()
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT
                COALESCE(SUM(vi."VendorInvoiceAmount"::numeric), 0) AS "TotalOwed",
                COUNT(*) AS "InvoiceCount"
            FROM "VendorInvoice" vi
            JOIN "PurchaseOrder" po ON vi."PurchaseOrderID" = po."PurchaseOrderID"
            WHERE COALESCE(vi."VendorInvoiceStatus", '') IN ('Draft', 'ReadyToPay')
              AND EXISTS (
                    SELECT 1
                    FROM "PurchaseOrderItem" poi
                    WHERE poi."PurchaseOrderID" = po."PurchaseOrderID"
                      AND COALESCE(poi."QuantityReceived", 0) > 0
              )
            '''
        )
        row = cur.fetchone() or {}
        return {
            "total_owed": float(row.get("TotalOwed") or 0),
            "invoice_count": int(row.get("InvoiceCount") or 0),
        }
    finally:
        conn.close()


def sweep_for_ready_to_pay_vendor_invoices():
    """LEGACY_AUTOMATION_DISABLED_BY_DEFAULT: runs the old ready-to-pay vendor invoice reminder sweeper only when explicitly enabled."""
    if not require_legacy_automation_runtime("database.vendor_invoices.sweep_for_ready_to_pay_vendor_invoices"):
        return

    from neon_ai.gateway import MY_EMAIL, send_to_user

    today = datetime.date.today()
    if today < AUTOMATION_MIN_DATE:
        return

    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute(
            '''
            SELECT vi."VendorInvoiceID", vi."VendorInvoiceNumber", vi."VendorInvoiceDate", vi."VendorInvoiceAmount", vi."VendorInvoiceStatus", po."PurchaseOrderID"
            FROM "VendorInvoice" vi
            JOIN "PurchaseOrder" po ON vi."PurchaseOrderID" = po."PurchaseOrderID"
            WHERE vi."VendorInvoiceStatus" = 'ReadyToPay'
              AND po."Date" >= %s
            ''',
            (AUTOMATION_MIN_DATE,),
        )
        rows = cur.fetchall()
    finally:
        conn.close()

    for row in rows:
        invoice_date = _parse_invoice_date(str(row.get("VendorInvoiceDate") or ""), reference_date=today)
        if not invoice_date or today < invoice_date + datetime.timedelta(days=25):
            continue

        memory = get_vendor_invoice_memory(row["VendorInvoiceID"])
        last_sent = memory.get("last_payment_reminder_at")
        if last_sent:
            try:
                last_date = datetime.datetime.fromisoformat(last_sent).date()
                if today <= last_date + datetime.timedelta(days=7):
                    continue
            except ValueError:
                pass

        report_path = generate_vendor_invoice_report(row["VendorInvoiceID"], force=False)
        send_to_user(
            subject=f"Payment Reminder: Vendor Invoice {row['VendorInvoiceNumber']}",
            content=(
                f"Vendor invoice {row['VendorInvoiceNumber']} for PO #{row['PurchaseOrderID']} has been ready to pay for 25 days.\n\n"
                f"Amount: ${float(row.get('VendorInvoiceAmount') or 0):,.2f}\n"
                f"Invoice date: {row.get('VendorInvoiceDate')}\n\n"
                "Please schedule payment."
            ),
            recipient=MY_EMAIL,
            attachment_path=report_path,
        )
        upsert_vendor_invoice_memory(row["VendorInvoiceID"], {"last_payment_reminder_at": datetime.datetime.now().isoformat()})
