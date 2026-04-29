from __future__ import annotations

import datetime
import os

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QComboBox,
    QDialog,
    QFrame,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPlainTextEdit,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)

from neon_ai.database.employees import get_all_employees
from neon_ai.database.timesheets import (
    add_manual_time,
    get_employee_timesheet,
    get_open_workorder_choices,
    get_standard_tasks,
    lock_week_timesheet,
)


class TimesheetManagerPage(QWidget):
    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.current_emp_id: int | None = None
        self.week_dates: list[datetime.date] = []
        self.emp_rate_dict: dict[int, float] = {}
        self.emp_record_map: dict[int, dict] = {}
        self.wo_choice_map: dict[str, int] = {}
        self.export_root_path = self.resolve_export_root()

        self._build_ui()
        self.load_employees()
        self.populate_weeks()

    def _build_ui(self) -> None:
        layout = QVBoxLayout(self)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.setSpacing(10)

        title = QLabel("Weekly Timesheet & Payroll Matrix")
        title.setStyleSheet("font-size: 20px; font-weight: 700;")
        layout.addWidget(title, 0, Qt.AlignmentFlag.AlignLeft)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        layout.addWidget(splitter, 1)

        left_group = QGroupBox("Staff Directory")
        left_layout = QVBoxLayout(left_group)
        self.emp_table = QTableWidget(0, 2)
        self.emp_table.setHorizontalHeaderLabels(["Emp #", "Name"])
        self.emp_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.emp_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.emp_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.emp_table.verticalHeader().setVisible(False)
        self.emp_table.setColumnWidth(0, 50)
        self.emp_table.horizontalHeader().setStretchLastSection(True)
        self.emp_table.itemSelectionChanged.connect(self.on_emp_select)
        left_layout.addWidget(self.emp_table)
        splitter.addWidget(left_group)

        self.right_frame = QWidget()
        right_layout = QVBoxLayout(self.right_frame)
        right_layout.setContentsMargins(10, 10, 10, 10)
        splitter.addWidget(self.right_frame)
        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 4)

        header = QFrame()
        header_layout = QHBoxLayout(header)
        header_layout.setContentsMargins(0, 0, 0, 0)
        self.lbl_emp_name = QLabel("Select an Employee")
        self.lbl_emp_name.setStyleSheet("font-size: 16px; font-weight: 700; color: #2980b9;")
        header_layout.addWidget(self.lbl_emp_name)
        header_layout.addStretch(1)
        header_layout.addWidget(QLabel("Select Week: "))
        self.week_combo = QComboBox()
        self.week_combo.setMinimumWidth(220)
        self.week_combo.currentTextChanged.connect(self.on_week_select)
        header_layout.addWidget(self.week_combo)
        right_layout.addWidget(header)

        matrix_group = QGroupBox("Weekly Log")
        matrix_layout = QVBoxLayout(matrix_group)
        self.ts_tree = QTreeWidget()
        self.ts_tree.setColumnCount(6)
        self.ts_tree.setHeaderLabels(["Day / Date", "WO #", "Task Description", "Source", "Status", "Hours"])
        self.ts_tree.setColumnWidth(0, 150)
        self.ts_tree.setColumnWidth(1, 80)
        self.ts_tree.setColumnWidth(3, 80)
        self.ts_tree.setColumnWidth(4, 80)
        self.ts_tree.setColumnWidth(5, 80)
        matrix_layout.addWidget(self.ts_tree)
        right_layout.addWidget(matrix_group, 1)

        summary_group = QGroupBox("Week Totals (Standard 8/40 Rule)")
        summary_layout = QHBoxLayout(summary_group)
        self.lbl_reg = QLabel("Regular Hours: 0.00")
        self.lbl_reg.setStyleSheet("font-size: 11px; font-weight: 700;")
        summary_layout.addWidget(self.lbl_reg)
        self.lbl_ot = QLabel("Overtime Hours: 0.00")
        self.lbl_ot.setStyleSheet("font-size: 11px; font-weight: 700; color: #c0392b;")
        summary_layout.addWidget(self.lbl_ot)
        self.lbl_total = QLabel("Total Hours: 0.00")
        self.lbl_total.setStyleSheet("font-size: 11px; font-weight: 700;")
        summary_layout.addWidget(self.lbl_total)
        summary_layout.addStretch(1)
        right_layout.addWidget(summary_group)

        entry_group = QGroupBox("Manual Entry (Appended to Selected Week)")
        entry_layout = QGridLayout(entry_group)
        entry_layout.addWidget(QLabel("Date:"), 0, 0)
        self.entry_date = QComboBox()
        self.entry_date.setMinimumWidth(110)
        entry_layout.addWidget(self.entry_date, 0, 1)
        entry_layout.addWidget(QLabel("WO #:"), 0, 2)
        self.entry_wo = QComboBox()
        self.entry_wo.setMinimumWidth(200)
        entry_layout.addWidget(self.entry_wo, 0, 3)
        entry_layout.addWidget(QLabel("Task:"), 0, 4)
        self.entry_task = QComboBox()
        self.entry_task.setMinimumWidth(180)
        entry_layout.addWidget(self.entry_task, 0, 5)
        entry_layout.addWidget(QLabel("Hours:"), 0, 6)
        self.entry_hours = QLineEdit()
        self.entry_hours.setMaximumWidth(60)
        entry_layout.addWidget(self.entry_hours, 0, 7)
        add_button = QPushButton("Add")
        add_button.clicked.connect(self.on_add_manual)
        entry_layout.addWidget(add_button, 0, 8)
        right_layout.addWidget(entry_group)

        footer = QFrame()
        footer_layout = QHBoxLayout(footer)
        footer_layout.setContentsMargins(0, 0, 0, 0)
        self.export_target_label = QLabel(f"Folder Target: {self.export_root_path}")
        self.export_target_label.setStyleSheet("color: gray;")
        footer_layout.addWidget(self.export_target_label)
        self.last_saved_label = QLabel("Last File: not saved yet")
        self.last_saved_label.setStyleSheet("color: #2980b9;")
        footer_layout.addWidget(self.last_saved_label)
        footer_layout.addStretch(1)

        send_employee_button = QPushButton("Send to Employee for approval")
        send_employee_button.clicked.connect(self.on_send_to_employee)
        footer_layout.addWidget(send_employee_button)
        send_folder_button = QPushButton("Send to Folder")
        send_folder_button.clicked.connect(self.on_send_to_folder)
        footer_layout.addWidget(send_folder_button)
        view_button = QPushButton("View")
        view_button.clicked.connect(self.on_view_timesheet)
        footer_layout.addWidget(view_button)
        lock_button = QPushButton("Lock")
        lock_button.clicked.connect(self.on_lock_week)
        footer_layout.addWidget(lock_button)
        save_button = QPushButton("Save")
        save_button.clicked.connect(self.on_save_timesheet)
        footer_layout.addWidget(save_button)
        right_layout.addWidget(footer)

    def resolve_export_root(self):
        preferred = r"D:\timesheets"
        fallback = r"C:\timesheets"
        root = preferred if os.path.exists("D:") else fallback
        os.makedirs(root, exist_ok=True)
        return root

    def load_employees(self) -> None:
        self.emp_table.setRowCount(0)
        self.emp_rate_dict = {}
        self.emp_record_map = {}
        try:
            for emp in get_all_employees():
                emp_id = int(emp["EmployeeID"])
                row = self.emp_table.rowCount()
                self.emp_table.insertRow(row)
                self.emp_table.setItem(row, 0, QTableWidgetItem(str(emp_id)))
                self.emp_table.setItem(row, 1, QTableWidgetItem(str(emp["EmployeeName"])))
                self.emp_rate_dict[emp_id] = float(emp.get("EmployeeRate") or 0)
                self.emp_record_map[emp_id] = emp
        except Exception as exc:
            print(f"Employee Load Error: {exc}")
        self.refresh_manual_entry_choices()

    def populate_weeks(self) -> None:
        today = datetime.date.today()
        year, week_num, _ = today.isocalendar()
        options = []
        for week in range(week_num, 0, -1):
            monday = datetime.date.fromisocalendar(year, week, 1)
            sunday = monday + datetime.timedelta(days=6)
            options.append(f"Wk {week}: {monday.strftime('%b %d')} - {sunday.strftime('%b %d')}, {year}")
        self.week_combo.blockSignals(True)
        self.week_combo.clear()
        self.week_combo.addItems(options)
        if options:
            self.week_combo.setCurrentIndex(0)
        self.week_combo.blockSignals(False)
        self.refresh_manual_entry_choices()

    def on_emp_select(self) -> None:
        selection = self.emp_table.selectedItems()
        if not selection:
            return
        row = self.emp_table.row(selection[0])
        self.current_emp_id = int(self.emp_table.item(row, 0).text())
        self.lbl_emp_name.setText(f"{self.emp_table.item(row, 1).text()}'s Timesheet")
        self.refresh_matrix()

    def on_week_select(self, _value: str | None = None) -> None:
        self.refresh_matrix()

    def refresh_matrix(self) -> None:
        if not self.current_emp_id or not self.week_combo.currentText():
            return

        week_str = self.week_combo.currentText()
        year = int(week_str.split(", ")[1])
        week_num = int(week_str.split(":")[0].replace("Wk ", ""))
        start_date = datetime.date.fromisocalendar(year, week_num, 1)
        self.week_dates = [start_date + datetime.timedelta(days=i) for i in range(7)]
        self.entry_date.clear()
        self.entry_date.addItems([day.isoformat() for day in self.week_dates])
        if self.entry_date.count():
            self.entry_date.setCurrentIndex(0)
        self.refresh_manual_entry_choices()

        self.ts_tree.clear()
        try:
            records = get_employee_timesheet(self.current_emp_id, self.week_dates[0], self.week_dates[-1])
            daily_totals = {day: 0.0 for day in self.week_dates}
            for row in records:
                day = row["DateWorked"]
                daily_totals[day] = daily_totals.get(day, 0.0) + float(row.get("HoursWorked") or 0)

            for day in self.week_dates:
                day_records = [row for row in records if row["DateWorked"] == day]
                parent = QTreeWidgetItem(
                    [
                        f"{day.strftime('%A (%b %d)')}",
                        "",
                        "",
                        "",
                        "",
                        f"{daily_totals.get(day, 0.0):.2f} hrs",
                    ]
                )
                self.ts_tree.addTopLevelItem(parent)
                parent.setExpanded(True)
                for row in day_records:
                    child = QTreeWidgetItem(
                        [
                            "",
                            str(row.get("WorkOrderID") or ""),
                            str(row.get("TaskName") or row.get("TaskId") or ""),
                            str(row.get("EntrySource") or ""),
                            str(row.get("Status") or ""),
                            f"{float(row.get('HoursWorked') or 0):.2f}",
                        ]
                    )
                    parent.addChild(child)

            regular_hours, overtime_hours = self.calculate_week_totals(daily_totals)
            total_hours = regular_hours + overtime_hours
            self.lbl_reg.setText(f"Regular Hours: {regular_hours:.2f}")
            self.lbl_ot.setText(f"Overtime Hours: {overtime_hours:.2f}")
            self.lbl_total.setText(f"Total Hours: {total_hours:.2f}")

            is_locked = bool(records) and all(str(row.get("Status") or "") in {"Approved", "Exported"} for row in records)
            if is_locked:
                self.lbl_total.setStyleSheet("font-size: 11px; font-weight: 700; color: green;")
                self.lbl_total.setText(f"Total Hours: {total_hours:.2f} (LOCKED)")
            else:
                self.lbl_total.setStyleSheet("font-size: 11px; font-weight: 700; color: black;")
        except Exception as exc:
            print(f"Matrix Refresh Error: {exc}")

    def calculate_week_totals(self, daily_totals):
        week_regular = 0.0
        week_overtime = 0.0
        for day in self.week_dates:
            hours = float(daily_totals.get(day, 0.0))
            if hours > 8:
                week_regular += 8
                week_overtime += hours - 8
            else:
                week_regular += hours
        if week_regular > 40:
            extra_ot = week_regular - 40
            week_overtime += extra_ot
            week_regular = 40
        return week_regular, week_overtime

    def refresh_manual_entry_choices(self) -> None:
        try:
            workorders = get_open_workorder_choices()
            self.wo_choice_map = {row["Label"]: int(row["WorkOrderID"]) for row in workorders}
            self.entry_wo.clear()
            self.entry_wo.addItems(self.wo_choice_map.keys())
            tasks = get_standard_tasks()
            self.entry_task.clear()
            self.entry_task.addItems(tasks)
        except Exception as exc:
            print(f"Manual Entry Choice Error: {exc}")

    def on_add_manual(self) -> None:
        if not self.current_emp_id:
            return
        try:
            selected_wo = self.entry_wo.currentText()
            wo_id = self.wo_choice_map.get(selected_wo)
            task_name = self.entry_task.currentText().strip()
            if not wo_id or not task_name:
                QMessageBox.critical(self, "Error", "Please choose a valid work order and task.")
                return

            data = {
                "emp_id": self.current_emp_id,
                "wo_id": int(wo_id),
                "date": self.entry_date.currentText(),
                "hours": float(self.entry_hours.text()),
                "task": task_name,
                "rate": self.emp_rate_dict.get(self.current_emp_id, 0.0),
            }
            if add_manual_time(data):
                self.entry_hours.setText("")
                if self.entry_task.count():
                    self.entry_task.setCurrentIndex(0)
                self.refresh_matrix()
        except ValueError:
            QMessageBox.critical(self, "Error", "Hours must be a number.")

    def _get_current_employee(self):
        if not self.current_emp_id:
            raise ValueError("Select an employee first.")
        employee = self.emp_record_map.get(self.current_emp_id)
        if not employee:
            raise ValueError("The selected employee record is not available.")
        return employee

    def _get_week_snapshot(self):
        if not self.current_emp_id or not self.week_dates:
            raise ValueError("Select both an employee and a week first.")

        employee = self._get_current_employee()
        records = get_employee_timesheet(self.current_emp_id, self.week_dates[0], self.week_dates[-1])
        daily_totals = {day: 0.0 for day in self.week_dates}
        for row in records:
            daily_totals[row["DateWorked"]] = daily_totals.get(row["DateWorked"], 0.0) + float(row.get("HoursWorked") or 0)
        regular_hours, overtime_hours = self.calculate_week_totals(daily_totals)
        return {
            "employee": employee,
            "records": records,
            "daily_totals": daily_totals,
            "regular_hours": regular_hours,
            "overtime_hours": overtime_hours,
            "total_hours": regular_hours + overtime_hours,
            "week_start": self.week_dates[0],
            "week_end": self.week_dates[-1],
            "is_locked": bool(records) and all(str(row.get("Status") or "") in {"Approved", "Exported"} for row in records),
        }

    def _safe_name(self, value):
        cleaned = "".join(ch if ch.isalnum() or ch in (" ", "-", "_") else "_" for ch in str(value or ""))
        return "_".join(cleaned.split()).strip("_") or "Timesheet"

    def _build_preview_text(self, snapshot):
        employee = snapshot["employee"]
        lines = [
            "WEEKLY TIMESHEET PREVIEW",
            "",
            f"Employee: {employee.get('EmployeeName') or 'Unknown'}",
            f"Email: {employee.get('EmployeeEmail') or 'No email on file'}",
            f"Week: {snapshot['week_start']} to {snapshot['week_end']}",
            f"Status: {'Locked' if snapshot['is_locked'] else 'Draft'}",
            "",
        ]
        for day in self.week_dates:
            lines.append(f"{day.strftime('%A %Y-%m-%d')} | Total: {snapshot['daily_totals'].get(day, 0.0):.2f}")
            day_rows = [row for row in snapshot["records"] if row["DateWorked"] == day]
            if not day_rows:
                lines.append("  No time entered")
            for row in day_rows:
                lines.append(
                    f"  WO #{row.get('WorkOrderID') or ''} | {row.get('TaskName') or row.get('TaskId') or ''} | "
                    f"{float(row.get('HoursWorked') or 0):.2f} hrs | {row.get('EntrySource') or ''} | {row.get('Status') or ''}"
                )
            lines.append("")
        lines.append(f"Regular Hours: {snapshot['regular_hours']:.2f}")
        lines.append(f"Overtime Hours: {snapshot['overtime_hours']:.2f}")
        lines.append(f"Total Hours: {snapshot['total_hours']:.2f}")
        lines.append(f"Folder Target: {self.export_root_path}")
        return "\n".join(lines)

    def _export_timesheet_docx(self, snapshot, label="Draft"):
        from docx import Document

        employee = snapshot["employee"]
        doc = Document()
        employee_name = employee.get("EmployeeName") or "Employee"
        doc.add_heading(f"Weekly Timesheet - {employee_name}", 0)
        doc.add_paragraph(f"Week: {snapshot['week_start']} to {snapshot['week_end']}")
        doc.add_paragraph(f"Document Status: {label}")
        doc.add_paragraph(f"Employee Email: {employee.get('EmployeeEmail') or 'No email on file'}")

        for day in self.week_dates:
            doc.add_heading(day.strftime("%A %b %d, %Y"), level=1)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "WO #"
            hdr[1].text = "Task"
            hdr[2].text = "Source"
            hdr[3].text = "Status"
            hdr[4].text = "Hours"

            day_rows = [row for row in snapshot["records"] if row["DateWorked"] == day]
            if not day_rows:
                row = table.add_row().cells
                row[0].text = ""
                row[1].text = "No time entered"
                row[2].text = ""
                row[3].text = ""
                row[4].text = "0.00"
            else:
                for entry in day_rows:
                    row = table.add_row().cells
                    row[0].text = str(entry.get("WorkOrderID") or "")
                    row[1].text = str(entry.get("TaskName") or entry.get("TaskId") or "")
                    row[2].text = str(entry.get("EntrySource") or "")
                    row[3].text = str(entry.get("Status") or "")
                    row[4].text = f"{float(entry.get('HoursWorked') or 0):.2f}"

            doc.add_paragraph(f"Daily Total: {snapshot['daily_totals'].get(day, 0.0):.2f} hours")

        doc.add_paragraph(f"Regular Hours: {snapshot['regular_hours']:.2f}")
        doc.add_paragraph(f"Overtime Hours: {snapshot['overtime_hours']:.2f}")
        doc.add_paragraph(f"Total Hours: {snapshot['total_hours']:.2f}")

        safe_employee = self._safe_name(employee_name)
        filename = f"Timesheet_{safe_employee}_{snapshot['week_start']}_{label}.docx"
        save_path = os.path.join(self.export_root_path, filename)
        doc.save(save_path)
        self.last_saved_label.setText(f"Last File: {save_path}")
        return save_path

    def on_save_timesheet(self) -> None:
        try:
            snapshot = self._get_week_snapshot()
            path = self._export_timesheet_docx(snapshot, label="Draft")
            QMessageBox.information(self, "Saved", f"Draft timesheet saved to:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Save Error", str(exc))

    def on_lock_week(self) -> None:
        try:
            snapshot = self._get_week_snapshot()
        except Exception as exc:
            QMessageBox.critical(self, "Lock Error", str(exc))
            return

        confirmed = QMessageBox.question(self, "Confirm Lock", "Approve all time for this week? It will be locked for payroll.")
        if confirmed != QMessageBox.StandardButton.Yes:
            return

        try:
            lock_week_timesheet(self.current_emp_id, snapshot["week_start"], snapshot["week_end"])
            self.refresh_matrix()
            snapshot = self._get_week_snapshot()
            path = self._export_timesheet_docx(snapshot, label="Locked")
            QMessageBox.information(self, "Locked", f"Week locked.\nTimesheet saved to:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Lock Error", str(exc))

    def on_view_timesheet(self) -> None:
        try:
            snapshot = self._get_week_snapshot()
            label = "Locked" if snapshot["is_locked"] else "Draft"
            path = self._export_timesheet_docx(snapshot, label=label)
            try:
                os.startfile(path)
                return
            except Exception:
                preview = QDialog(self)
                preview.setWindowTitle("Timesheet Preview")
                preview.resize(900, 700)
                layout = QVBoxLayout(preview)
                text = QPlainTextEdit()
                text.setReadOnly(True)
                text.setPlainText(self._build_preview_text(snapshot))
                layout.addWidget(text)
                preview.exec()
        except Exception as exc:
            QMessageBox.critical(self, "View Error", str(exc))

    def on_send_to_folder(self) -> None:
        try:
            snapshot = self._get_week_snapshot()
            label = "Locked" if snapshot["is_locked"] else "Draft"
            path = self._export_timesheet_docx(snapshot, label=label)
            try:
                os.startfile(self.export_root_path)
            except Exception:
                pass
            QMessageBox.information(self, "Exported", f"Timesheet exported to:\n{path}\n\nFolder:\n{self.export_root_path}")
        except Exception as exc:
            QMessageBox.critical(self, "Export Error", str(exc))

    def on_send_to_employee(self) -> None:
        try:
            snapshot = self._get_week_snapshot()
            employee = snapshot["employee"]
            recipient = str(employee.get("EmployeeEmail") or "").strip()
            if not recipient:
                raise ValueError("This employee does not have an email address on file.")

            label = "Locked" if snapshot["is_locked"] else "Draft"
            path = self._export_timesheet_docx(snapshot, label=label)

            from neon_ai.gateway import MY_EMAIL, send_to_user

            sent = send_to_user(
                subject=f"Timesheet Review Needed - {employee.get('EmployeeName')} - Week of {snapshot['week_start']}",
                content=(
                    f"Hello {employee.get('EmployeeName')},\n\n"
                    f"Please review your timesheet for the week of {snapshot['week_start']} to {snapshot['week_end']}.\n\n"
                    "If anything is incorrect, reply with the required corrections before payroll processing.\n\n"
                    "Thank you,\nArgon Electrical"
                ),
                recipient=recipient,
                attachment_path=path,
                cc_recipients=[MY_EMAIL],
            )
            if not sent:
                raise RuntimeError("The email gateway could not send the timesheet.")
            QMessageBox.information(self, "Sent", f"Timesheet sent to {recipient}.\n\nAttachment:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Email Error", str(exc))

    def refresh_data(self) -> None:
        self.load_employees()
        self.populate_weeks()
        self.ts_tree.clear()
        self.lbl_emp_name.setText("Select an Employee")
        self.lbl_reg.setText("Regular Hours: 0.00")
        self.lbl_ot.setText("Overtime Hours: 0.00")
        self.lbl_total.setText("Total Hours: 0.00")
        self.lbl_total.setStyleSheet("font-size: 11px; font-weight: 700;")
        self.entry_hours.setText("")
