import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_workorder_docx(data, filepath):
    doc = docx.Document()
    
    # Header
    header = doc.add_heading('FIELD JOB PACKET', 0)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Job Details Table
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    
    table.cell(0,0).text = f"WORK ORDER #: {data['WorkOrderID']}"
    table.cell(0,1).text = f"SOURCE EST #: {data['EstimateID'] or 'N/A'}"
    
    table.cell(1,0).text = f"CUSTOMER: {data['CustomerName']}"
    table.cell(1,1).text = f"DATE ISSUED: {data['Date']}"
    
    street_info = f"{data.get('StreetNumber', '')} {data.get('StreetName', '')}".strip()
    table.cell(2,0).text = f"SITE ADDRESS: {street_info}, {data.get('City', 'N/A')}"
    table.cell(2,1).text = f"CUSTOMER PO: {data['CustomerPO']}"
    
    table.cell(3,0).text = f"CONTACT PHONE: {data['Phone']}"
    table.cell(3,1).text = "STATUS: OPEN / ACTIVE"

    # Full Scope
    doc.add_heading('Scope of Work', level=1)
    doc.add_paragraph(data['WorkOrderScope'] or "Refer to field blueprints.")

    # Anticipated Resources
    doc.add_heading('Anticipated Labor & Materials', level=1)
    doc.add_heading('Estimated Tasks:', level=2)
    doc.add_paragraph(data['EstLaborSummary'] or "No specific tasks listed in estimate.")
    
    doc.add_heading('Materials Required:', level=2)
    doc.add_paragraph(data['EstMaterialSummary'] or "Standard material requirements.")

    doc.add_page_break()
    
    # Installer Logs
    doc.add_heading('Installer Field Logs', level=1)
    doc.add_paragraph("Please record actual time and materials used on site.")
    doc.add_table(rows=8, cols=4).style = 'Table Grid' # Labor Log

    doc.save(filepath)