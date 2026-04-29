import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def _safe_filename(text):
    return "".join(c for c in text if c not in '\\/*?:"<>|').strip().replace(" ", "_")

def export_estimate_doc(d, destination_dir=None, filename=None):
    doc = Document()
    
    # USE THIS STYLE - It is built into Word by default
    SAFE_STYLE = 'Table Grid' 

    # When creating tables, apply it immediately:
    table = doc.add_table(rows=1, cols=2)
    table.style = SAFE_STYLE
    # --- HEADER ---
    title = doc.add_heading('PROPOSED ESTIMATE', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- CONTACT INFO TABLE (2 Columns) ---
    table_info = doc.add_table(rows=1, cols=2)
    table_info.width = Inches(6.0)
    
    # Left Side: Customer
    cust_cell = table_info.rows[0].cells[0]
    cust_cell.text = f"BILL TO:\n{d['customer_name']}\n{d['customer_info']}"
    
    # Right Side: Site
    site_cell = table_info.rows[0].cells[1]
    site_cell.text = f"PROJECT SITE:\n{d['site_name']}\n{d['site_addr']}\n\nDate: {d['date']}\nEst ID: {d['id']}"

    # --- SCOPE OF WORK ---
    doc.add_heading('Scope of Work', level=1)
    doc.add_paragraph(d['scope'])

    # --- LABOR LIST ---
    doc.add_heading('Estimated Labor Services', level=2)
    lab_table = doc.add_table(rows=1, cols=2)
    lab_table.style = SAFE_STYLE
    hdr_cells = lab_table.rows[0].cells
    hdr_cells[0].text = 'Role / Service'
    hdr_cells[1].text = 'Amount'
    
    lab_subtotal = 0
    for item in d['labor_items']:
        row = lab_table.add_row().cells
        row[0].text = item['role']
        row[1].text = f"${item['total']:,.2f}"
        lab_subtotal += item['total']

    # --- MATERIAL LIST ---
    doc.add_heading('Estimated Material Requirements', level=2)
    mat_table = doc.add_table(rows=1, cols=2)
    mat_table.style = SAFE_STYLE
    hdr_cells = mat_table.rows[0].cells
    hdr_cells[0].text = 'Description'
    hdr_cells[1].text = 'Amount'
    
    mat_subtotal = 0
    for item in d['mat_items']:
        row = mat_table.add_row().cells
        row[0].text = item['desc']
        row[1].text = f"${item['total']:,.2f}"
        mat_subtotal += item['total']

    # --- FINAL TOTAL ---
    doc.add_paragraph("\n")
    total_table = doc.add_table(rows=3, cols=2)
    
    # Labor Sub
    row = total_table.rows[0].cells
    row[0].text = "Labor Subtotal:"
    row[1].text = f"${lab_subtotal:,.2f}"
    
    # Material Sub
    row = total_table.rows[1].cells
    row[0].text = "Material Subtotal:"
    row[1].text = f"${mat_subtotal:,.2f}"
    
    # GRAND TOTAL
    row = total_table.rows[2].cells
    row[0].text = "TOTAL ESTIMATED PROJECT PRICE:"
    row[1].text = f"${(lab_subtotal + mat_subtotal):,.2f}"
    row[0].paragraphs[0].runs[0].font.bold = True
    row[1].paragraphs[0].runs[0].font.bold = True

    # --- SAVE ---
    if destination_dir is None:
        destination_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "OutBox")
    os.makedirs(destination_dir, exist_ok=True)

    if filename is None:
        filename = f"Detailed_Estimate_{d['id']}_{_safe_filename(d['customer_name'])}.docx"

    final_path = os.path.join(destination_dir, filename)
    doc.save(final_path)
    return final_path

def export_to_outbox(d):
    return export_estimate_doc(d)
